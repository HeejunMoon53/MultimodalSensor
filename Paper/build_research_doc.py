"""연구계획 + 논문 스케치 단일 문서 생성.

python Paper/build_research_doc.py  →  Paper/연구계획_논문스케치.docx
  1부: Docs/plan/연구계획.md 를 그대로 옮긴 계획 (원본은 md)
  2부: 논문 스케치 — 빈칸은 노란 바탕 [G-xx], 보유 그림은 세미나 원본(Paper/sketch/seminar_images)·
       슬라이드(Paper/sketch/seminar_slides)·시뮬(sim_gripper/out/paper_figs)
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PLAN_MD = ROOT / "Docs" / "plan" / "연구계획.md"
SLIDES = HERE / "sketch" / "seminar_slides"
IMGS = HERE / "sketch" / "seminar_images"
SIM = ROOT / "sim_gripper" / "out" / "paper_figs"
OUT = HERE / "연구계획_논문스케치.docx"
GREEN, PURPLE = RGBColor(0x1B, 0x6E, 0x3A), RGBColor(0x5B, 0x3E, 0x9E)


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "맑은 고딕", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        d.styles[s].font.name = "맑은 고딕"
        d.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        d.styles[s].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
    return d


# ── 1부: markdown → docx ────────────────────────────────────────────────
def _inline(p, text):
    """**굵게**, `코드`만 처리."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            p.add_run(part)


def render_md(d, md):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            p = d.add_paragraph()
            r = p.add_run("\n".join(lines[i + 1:j]))
            r.font.name, r.font.size = "Consolas", Pt(8.5)
            i = j + 1
            continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = d.add_table(rows=0, cols=len(rows[0]))
            t.style = "Light Grid Accent 1"
            for k, r in enumerate(rows):
                cells = t.add_row().cells
                for c, x in zip(cells, r):
                    c.text = ""
                    _inline(c.paragraphs[0], x)
                    for run in c.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
                        if k == 0:
                            run.bold = True
            d.add_paragraph()
            continue
        m = re.match(r"^(#{1,3}) (.*)", ln)
        if m:
            level = len(m.group(1))
            if level == 1:
                d.add_heading(m.group(2), 0)
            else:
                d.add_heading(m.group(2), level - 1)
        elif ln.startswith("> "):
            p = d.add_paragraph()
            _inline(p, ln[2:])
            for r in p.runs:
                r.italic, r.font.color.rgb = True, RGBColor(0x55, 0x55, 0x55)
        elif re.match(r"^\s*- \[[ x]\] ", ln):
            done = "[x]" in ln
            p = d.add_paragraph(style="List Bullet")
            _inline(p, ("☑ " if done else "☐ ") + re.sub(r"^\s*- \[[ x]\] ", "", ln))
        elif re.match(r"^\s*[-*] ", ln):
            indent = len(ln) - len(ln.lstrip())
            p = d.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            _inline(p, re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\d+\. ", ln):
            p = d.add_paragraph(style="List Number")
            _inline(p, re.sub(r"^\d+\. ", "", ln))
        elif ln.strip():
            _inline(d.add_paragraph(), ln)
        i += 1


# ── 2부: 논문 스케치 ─────────────────────────────────────────────────────
class Sketch:
    def __init__(self, d):
        self.d = d

    def para(self, *runs, style=None, center=False):
        p = self.d.add_paragraph(style=style)
        for r in runs:
            text, kind = (r, None) if isinstance(r, str) else r
            run = p.add_run(text)
            if kind == "blank":
                run.font.highlight_color, run.bold = WD_COLOR_INDEX.YELLOW, True
            elif kind == "have":
                run.font.color.rgb, run.bold = GREEN, True
            elif kind == "sim":
                run.font.color.rgb, run.bold = PURPLE, True
            elif kind == "b":
                run.bold = True
            elif kind == "i":
                run.italic = True
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    @staticmethod
    def blank(text):
        return (f"⬜ {text}", "blank")

    def bullet(self, *runs):
        return self.para(*runs, style="List Bullet")

    def claim(self, text, status):
        self.para(("이 절의 주장: ", "b"), text, "   ", ("근거: ", "b"), *status)

    def slide(self, n, label, width=11.5):
        self._pic([SLIDES / f"S{n:02d}.png"], width)
        self.para((f"▲ {label}  ", "i"), (f"[세미나 S{n}]", "have"), center=True)

    def imgs(self, names, label, width=7.5, note=None):
        files = [next(IMGS.glob(f"{n}.*"), None) for n in names]
        self._pic([f for f in files if f], width)
        self.para((f"▲ {label}  ", "i"), (f"[세미나 원본 {', '.join(names)}]", "have"), center=True)
        if note:
            self.para(self.blank(note), center=True)

    def sim(self, name, label, width=13):
        self._pic([SIM / f"{name}.png"], width)
        self.para((f"▲ {label}  ", "i"), (f"[시뮬 {name}]", "sim"), center=True)

    def slot(self, label, gap):
        self.para(self.blank(f"[그림 자리] {label}  ({gap})"), center=True)

    def _pic(self, paths, width):
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for f in paths:
            if Path(f).exists():
                p.add_run().add_picture(str(f), width=Cm(width))
                p.add_run("  ")


def add_sketch(d):
    S = Sketch(d)
    B = S.blank
    d.add_page_break()
    d.add_heading("2부. 논문 스케치 (틀)", 0)
    S.para(("표시: ", "b"), ("초록 = 보유(세미나)", "have"), "  ·  ", ("보라 = 보유(시뮬레이션)", "sim"), "  ·  ",
           B("노란 칸 = 비어 있음 [G-xx] (1부 7절)"), "  ·  모듈 M0–M5 = 1부 4절")
    S.para("문장은 다듬지 않은 스케치다. 숫자는 세미나 값이고, 빈 칸은 채우지 않고 남긴다.")

    d.add_heading("Abstract (스케치)", 1)
    S.para("Robots that physically interact with humans need to perceive their own deformation and nearby objects "
           "before and during contact. Existing soft multimodal sensors stack multiple sensing elements or rely on "
           "external instruments, increasing structural and electrical complexity. Here we present a single-electrode "
           "(two-wire) liquid-metal spiral coil that senses ",
           ("strain, proximity, contact force, contact events and the electrical property of nearby objects", "b"),
           " by time-division measurement of inductance, DC resistance and triboelectric voltage within a 1-ms cycle. "
           "A gated mixture-of-experts network with causal EMA features (", ("≈1.4 k parameters, 5.5 KB", "have"),
           ") decouples the coupled signals on an STM32 MCU in ", ("≈260 µs", "have"), ". On held-out data it achieves "
           "R² of ", ("0.988 (strain), 0.990 (proximity), 0.982 (force)", "have"), " with ",
           ("98.7 % contact classification", "have"), "; real-time on-board validation yields ", B("R² [G-E1]"),
           ". The sign of the inductance residual identifies conductive versus dielectric objects (",
           B("accuracy [G-E2]"), "), and the TENG channel detects contact and separation ", B("faster / more reliably "
                                                                                              "than the gate [G-E3]"),
           ". Finally, the estimated modalities drive a state-machine controller for a soft gripper — "
           "proximity-scheduled pre-shaping, material-adaptive force control and a safety layer — validated in "
           "simulation.")
    S.para(("핵심 메시지: ", "b"), "C1 다섯 모달리티 · C2 1 ms TDM · C3 MCU 디커플링 · C4 모달리티 기반 그리퍼 제어 (1부 1절)")

    # 1 Introduction
    d.add_heading("1. Introduction", 1)
    S.claim("단일 전극·온보드 AI로 구조·전기적 복잡성을 동시에 줄인다", [("배경·문제 보유", "have"), " / ", B("차별점 표 [G-W1]")])
    S.bullet("배경: HRI에서 물리적 지각(자기 변형 + 주변 물체), 소프트 센서의 장점 — S4–S5")
    S.bullet("문제: 구조적 복잡성(적층·배선·박리) + 전기적 복잡성(crosstalk·외부 계측·PC 의존) — S6")
    S.bullet("접근: 단일 평면 전극 + 시분할 다중 신호 + 온보드 경량 디커플링 + 온보드 제어 — S7–S8")
    S.bullet(("Related work ", None), B("다중 모달 소프트 센서 / 인덕티브·LC 센서 / TENG / 엣지 AI [G-W2]"))
    S.bullet(("노벨티 문장 ", None), B("기존 연구 대비 차별점 3개 — 세미나 P.6 비교 보강 [G-W9]"))
    S.bullet(("EGaIn을 쓰는 이유 ", None), B("신축성·전도도·코일 형상 유지·R/L 동시 응답, 대안 재료 비교 [G-W8]"))
    S.bullet("기여: C1–C4 (C4는 시뮬레이션 검증)")
    S.slide(7, "Fig. 1  Whole process of multimodal soft single-sensor system")
    S.slide(6, "(참고) 기존 다중 모달 센서의 문제 — Introduction 그림 후보", width=10)
    S.slot("Table 1  선행연구 비교 (모달리티 수·전극 수·계측 장비·디커플링·온보드 여부)", "G-W1 · M4")

    # 2 Sensor
    d.add_heading("2. Sensor Design, Fabrication and Working Principle", 1)
    S.claim("한 코일의 L·R·V가 자극마다 다른 방향으로 반응한다 → 분리 가능",
            [("자극–응답 보유", "have"), " / ", B("재질·TENG 정량, 반복성, 물리 해석 [G-E2·E3·E4·E12]")])
    S.bullet("구조: EGaIn DIW 평면 사각 스파이럴, 약 20 × 40 mm, 11턴, 선폭 ~400 µm, 피치 ~700 µm, 브리지 — S10")
    S.bullet(("돌기 구조로 압축 감도 향상 ", None), B("효과 데이터 [G-E5, 초록 이후]"))
    S.bullet("제작 6단계 — S11  ", B("공정 조건 수치 [G-W3]"))
    S.bullet("자극–응답: 인장 f↓R↑ / 압축 f↑R↑ / 도체 근접 f↑ / 유전체 근접 f↓ / 접촉 V↑ — S12")
    S.bullet("메커니즘: 길이·단면(R, L), 와전류(L↓), 기생 정전용량(L↑), 마찰대전(V) — S13–S14")
    S.bullet(("물리적 근거 ", None), B("이론식 ↔ 실측 [G-W6], 전자기·기계 FEA 경향 재현 [G-E12 · M3]"))
    S.slide(10, "Fig. 2a  Sensor design (dimensions, bumps, bridge)")
    S.imgs(["S10_01", "S10_06"], "Fig. 2b  Fabricated sensor (relaxed / stretched)", width=5.0)
    S.imgs(["S10_03", "S10_08"], "Fig. 2c  Microscope: channel cross-section / trace top view", width=6.5)
    S.slide(11, "Fig. 2d  Fabrication process")
    S.imgs(["S12_14"], "Fig. 3a  Stimulus–response (stretch · compression · conductor · dielectric · contact)", width=12,
           note="(a)–(d)가 실측인지 피팅인지 확인 [G-F1]")
    S.imgs(["S12_02"], "Fig. 3b  Proximity by material — hand / steel / acryl (재질 모달리티 예비 근거)", width=14,
           note="재질 확대·반복·분류 정량 [G-E2 · M1]")
    S.imgs(["S12_05"], "Fig. 3c  Contact (TENG) and elongation tests", width=14,
           note="속도·힘·재질별 TENG 정량 [G-E3 · M2]")
    S.imgs(["S14_19", "S12_19"], "Fig. 3d  Mechanisms: eddy current / parasitic capacitance & cross-section", width=7.0,
           note="S14_04·05 등 모식도 출처 확인 → 재작도 [G-F9]")
    S.slot("FEA: L(ε, d) 도체·유전체 의존성 — 측정 경향과 비교", "G-E12 · G-F10 · M3")
    S.slot("반복성·히스테리시스·센서 간 편차", "G-E4 · G-F7")

    # 3 Measurement
    d.add_heading("3. Measurement System", 1)
    S.claim("2선 하나에서 L·R·V를 1 ms 안에 시분할 측정한다", [("PCB·TDM·타이밍 보유", "have"), " / ", B("TDM 특성 정량 [G-E6, 초록 이후]")])
    S.bullet("스위칭 PCB (STM32G473 + LDC1614 + ADG734 MUX + OPAmp) — S16–S17")
    S.bullet("TDM: TIM7 1 ms, I²C-DMA(L)와 ADC-DMA(V 100 µs → R 100 µs) 병렬 — S18")
    S.bullet(("사양표 ", None), B("LDC 설정·공진 커패시터·MUX·ADC [G-W4]"))
    S.imgs(["S16_02", "S16_03"], "Fig. 4a  Switching PCB — schematic / 3D render", width=7.5)
    S.slide(17, "Fig. 4b  Single sensor, two wires → L / R / V by MUX")
    S.imgs(["S18_01", "S18_03"], "Fig. 4c  TDM: non-blocking DMA flow / acquisition timing", width=8.0)
    S.imgs(["S13_10"], "Fig. 4d  Channel switching waveform", width=8.0, note="전환 과도·crosstalk 정량 [G-E6]")
    S.sim("fig2_timing_budget", "Fig. 4e  1-ms cycle timing & model budget (제어기 시간은 추정 → G-E8)")

    # 4 Characterization
    d.add_heading("4. Signal Characterization and Dataset", 1)
    S.claim("비접촉에서 R은 인장만, L은 인장+근접에 의존 → 디커플링의 물리적 근거",
            [("플랫폼·프로토콜 A/B/C·곡면 보유", "have"), " / ", B("프로토콜 D(재질)·E(TENG) [G-E2·E3]")])
    S.bullet("5축 리니어 스테이지 (700 × 700 mm, 인장 4축 + 근접 1축), F/T 센서 정답 — S20–S21")
    S.bullet("프로토콜 A(인장 연속·근접 이산) / B(근접 연속·인장 이산) / C(인장 이산·힘 연속) — S22")
    S.bullet(("데이터셋 통계 ", None), B("세션·샘플·분할 [G-W5]"))
    S.bullet(("디커플링 가능성 정량 ", None), B("L·R 등고선 직교성·야코비안 조건수 지도 [G-E11 · M3]"))
    S.imgs(["S20_01", "S20_03"], "Fig. 5a  Custom 5-axis test platform — render / photo", width=7.5)
    S.slide(21, "Fig. 5b  Ground truth (F/T, stage) and sensor data synchronization")
    S.slide(22, "Fig. 5c  Data acquisition protocols A / B / C")
    S.imgs(["S23_01", "S23_02", "S23_03"], "Fig. 5d  Time series — A / B / C", width=5.3)
    S.imgs(["S24_04"], "Fig. 5e  Non-contact surfaces: ΔL(ε, d), ΔR(ε, d)", width=14)
    S.imgs(["S24_01"], "Fig. 5f  Contact surfaces: ΔL(ε, F), ΔR(ε, F)", width=12)

    # 5 Decoupling
    d.add_heading("5. Embedded Decoupling for Real-Time Inference", 1)
    S.claim("1.4 k 파라미터 Gate-MoE가 MCU에서 260 µs에 디커플링하고, 해석적 역산보다 안정적이다",
            [("구조·오프라인·배포 보유", "have"), " / ", B("실시간 성능 [M0], 선정 근거·기준선 [M3]")])
    S.bullet("SLS-EMA + Gate-MoE: 입력 (ΔL, ΔR, EMA_L, EMA_R), 게이트 → 비접촉(ε, d) / 접촉(ε, F), 전이 시 EMA 리셋 — S26")
    S.bullet(("오프라인: 게이트 0.987, ε R² 0.988 (0.108 %), d R² 0.990 (0.318 mm), F R² 0.982 (0.253 N)", "have"), " — S27")
    S.bullet(("배포: sklearn → ONNX → ST Edge AI → STM32 C, 1,416 파라미터, 5.53 KB, ≈260 µs (주기의 26 %)", "have"), " — S28")
    S.bullet(("실시간: 게이트 0.989, ε 0.947, d 0.794, F 0.591", "have"), " — S30  ", B("수정 후 재측정 [G-E1 · M0]"))
    S.bullet(("모델 선정 근거 ", None), B("단일 MLP vs MoE, EMA 유무·τ, 크기 스윕 → 정확도–파라미터–지연 파레토 [G-E11 · M3]"))
    S.bullet(("기준선 ", None), B("해석적 역산 vs Gate-MoE (실측 test) [G-E7 · M3]"))
    S.slide(26, "Fig. 6a  SLS-EMA + Gate-MoE architecture")
    S.imgs(["S27_02"], "Fig. 6b  Offline held-out parity", width=12)
    S.imgs(["S27_04"], "Fig. 6c  Expert-wise parity", width=14)
    S.slot("모델 선정 파레토 (정확도 vs 파라미터 vs 지연)", "G-E11 · G-F10 · M3")
    S.sim("fig7_analytic_vs_moe", "Fig. 6d  Analytical inversion vs MoE (시뮬 — 실측 test로 교체 [G-E7])")
    S.slide(28, "Fig. 7a  Embedded deployment pipeline & budget")
    S.slide(29, "Fig. 7b  Real-time validation setup")
    S.imgs(["S30_02"], "Fig. 7c  On-hardware real-time validation (separate session)", width=14,
           note="'smoothed, 3rd force peak removed' 처리 명시 또는 원데이터 재작도, 축 라벨 %→mm [G-F2]")
    S.slot("수정된 실시간 검증 (동기화·폐루프 후)", "G-E1 · M0")

    # 6 Material & contact
    d.add_heading("6. Electrical-Property (Material) and Contact Sensing", 1)
    S.claim("접촉 전에 재질을 알고, TENG로 접촉·분리를 게이트보다 빠르고 넓게 잡는다",
            [("예비(S12_02, 시뮬) 보유", "have"), " / ", B("정량 전부 [M1·M2]")])
    S.bullet("재질: 비접촉 ΔR로 역산한 ε로 ΔL의 변형률 성분을 빼고, 남은 잔차의 부호로 도체/유전체를 판별한다 (PINN Stage 1과 같은 원리)")
    S.bullet("접촉: 게이트 ∨ TENG 래치")
    S.bullet(("TENG가 필요한 이유 (피드백 1-1) ", None), B("검출 지연·유전체/저하중 검출률·분리 검출: 게이트만 vs 게이트 ∨ TENG [G-E3 · M2]"))
    S.bullet(("통합 모델 ", None), B("재질·TENG 포함 온보드 모델, 예산 갱신 [G-E9]"))
    S.sim("fig5_material_adaptive", "(예비) 재질 잔차와 재질별 목표 힘 — 시뮬")
    S.slot("재질 분류 혼동행렬 · 잔차 vs 거리 (실측)", "G-E2 · G-F3 · M1")
    S.slot("TENG 조건별 전압 · 검출 지연 · 게이트 대비 비교 (실측)", "G-E3 · G-F4 · M2")

    # 7 Application
    d.add_heading("7. Application: Soft Gripper with Perception and Adaptive Control", 1)
    S.claim("다섯 모달리티가 각각 제어 결정에 쓰인다 (시뮬레이션 검증)", [("시뮬 F1·F3·F4·F8 보유", "sim"), " / ", B("실물 [G-E10, 2027]")])
    S.bullet("상태 머신 (OPEN → APPROACH → GRASP → HOLD → RELEASE), 근접 스케줄 pre-shape, 재질 적응 admittance, 안전 계층 — 1 kHz 온보드")
    S.bullet("ablation: 재질을 빼면 계란 파손, TENG를 빼면 파지 실패, 근접을 빼면 충격 증가 (시뮬)")
    S.sim("fig1_architecture", "Fig. 8a  Sensing-to-control pipeline")
    S.sim("fig3_state_machine", "Fig. 8b  State machine & modality roles")
    S.sim("fig4_grasp_timeline", "Fig. 9  Five modalities in one grasp (simulated)")
    S.sim("fig8_ablation", "Fig. 10  Modality ablation (simulated)")
    S.slot("실물 그리퍼 파지 사진·로그", "G-E10 · G-F6 · M5 (2027)")

    # 8 Conclusion
    d.add_heading("8. Conclusion", 1)
    S.bullet("요약: S32 + 재질·TENG·제어 결과 ", B("[M1, M2]"))
    S.bullet("향후: 실물 그리퍼, 다중 주파수 임피던스(재질 인식 강화), 2D 변형, 무선 전력·통신 TDM, 어레이")

    d.add_heading("부록. 그림·표 계획", 1)
    rows = [("Fig. 1", "시스템 개요", "S7", "보유", "—"),
            ("Fig. 2", "센서 설계·실물·현미경·제작", "S10, S11 + 원본", "보유", "G-E5, G-W3"),
            ("Fig. 3", "자극–응답·재질 비교·TENG·메커니즘", "S12_14, S12_02, S12_05, S14_19", "부분", "G-F1, G-F9, G-E2, G-E3, G-E12"),
            ("Fig. 4", "PCB·TDM·전환 파형·타이밍", "S16–S18, S13_10 + 시뮬 F2", "보유", "G-E6, G-E8, G-W4"),
            ("Fig. 5", "플랫폼·동기화·프로토콜·시계열·곡면", "S20–S24", "보유 (D·E 없음)", "G-E2, G-E3, G-W5"),
            ("Fig. 6", "Gate-MoE·parity·선정 근거·기준선", "S26, S27 + 시뮬 F7", "부분", "G-E7, G-E11"),
            ("Fig. 7", "배포·실시간 검증", "S28–S30", "보유 (수정 필요)", "G-E1, G-F2"),
            ("신규", "재질 분류·TENG 정량", "—", "없음", "G-E2, G-E3, G-E9"),
            ("Fig. 8–10", "제어 구조·시뮬 결과", "시뮬 F1, F3, F4, F8", "보유 (시뮬)", "G-E10"),
            ("Table 1", "선행연구 비교", "—", "없음", "G-W1"),
            ("Table 2", "모델 예산·지연", "S28", "보유", "G-E8")]
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ("번호", "내용", "출처", "상태", "보충 ID")):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for c, x in zip(cells, r):
            c.text = x
        if r[3] == "없음":
            cells[3].paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def main():
    d = new_doc()
    render_md(d, PLAN_MD.read_text(encoding="utf-8"))
    add_sketch(d)
    try:
        d.save(str(OUT))
        print(OUT)
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_갱신.docx")
        d.save(str(alt))
        print(alt)


if __name__ == "__main__":
    main()
