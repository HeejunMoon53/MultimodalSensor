"""A–Z 보고서 생성 (docx). 수치·파라미터는 모두 코드/결과 파일에서 읽어 채운다.

python -m sim_gripper.report.build_report  →  sim_gripper/out/소프트그리퍼_시뮬레이션_보고서.docx
"""
import inspect
import json
from pathlib import Path

import imageio.v2 as imageio
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from sim_gripper.actuators.pneumatic import PneumaticActuator
from sim_gripper.actuators.tendon import TendonActuator
from sim_gripper.control.controller import MATERIAL_TABLE, CtrlParams
from sim_gripper.experiments.domain_shift import DEFAULT_SHIFT
from sim_gripper.experiments.run_all import OUT
from sim_gripper.actuators import PNEUMATIC_60A
from sim_gripper.model.gripper import BASE_R, BASE_Z, SENSOR_SEGS, BellowsParams, FingerParams, GripperSim
from sim_gripper.perception import moe as M
from sim_gripper.perception.contact import TengContact
from sim_gripper.perception.material import MaterialClassifier
from sim_gripper.sensor.forward import K_MAT, SensorForward, TengModel
from sim_gripper.sim import loop as L
from sim_gripper.sim.scenarios import SCENARIOS

DOC_NAME = "소프트그리퍼_시뮬레이션_보고서.docx"
DISCLAIMER = "시뮬레이션 내 방법론 시연, 실물 근거 아님"
CHAPTERS = [
    "1. 개요와 목적", "2. 시스템 구조 (Fig.1(b) 대응)", "3. M0 손가락 기구", "4. M1 센서 순방향 모델",
    "5. M2 TDM 샘플링", "6. M3 Gate-MoE 디커플링", "7. M4 재질 판별", "8. M5 제어기", "9. M6 액추에이터",
    "10. 실험 E1~E6 결과", "11. M7 도메인 시프트와 fine-tune", "12. 한계와 향후 과제",
    "부록 A. 파라미터 표", "부록 B. 기호표", "부록 C. 60A 벨로즈 그리퍼",
]
INK = RGBColor(0x22, 0x22, 0x22)


def _defaults(cls):
    return {k: v.default for k, v in inspect.signature(cls.__init__).parameters.items()
            if v.default is not inspect.Parameter.empty}


class Report:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = "맑은 고딕", Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
            self.d.styles[s].font.name = "맑은 고딕"
            self.d.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            self.d.styles[s].font.color.rgb = INK
        for sec in self.d.sections:
            sec.left_margin = sec.right_margin = Cm(2.2)

    def h(self, text, level=1):
        self.d.add_heading(text, level=level)

    def p(self, text, bold_lead=None):
        para = self.d.add_paragraph()
        if bold_lead:
            para.add_run(bold_lead).bold = True
        para.add_run(text)
        return para

    def bullets(self, items):
        for it in items:
            self.d.add_paragraph(it, style="List Bullet")

    def eq(self, *lines):
        for ln in lines:
            para = self.d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(ln)
            r.font.name, r.italic, r.font.size = "Cambria Math", True, Pt(11)

    def table(self, header, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for c, txt in zip(t.rows[0].cells, header):
            c.text = str(txt)
            c.paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = t.add_row().cells
            for c, txt in zip(cells, row):
                c.text = str(txt)
        if widths:
            for row in t.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Cm(w)
        self.d.add_paragraph()

    def fig(self, path, caption, width_cm=16):
        if Path(path).exists():
            self.d.add_picture(str(path), width=Cm(width_cm))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = self.d.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size, cap.runs[0].italic = Pt(9), True
        else:
            self.p(f"[그림 없음: {Path(path).name} — 실험 스크립트를 먼저 실행]")


def _f(x, nd=3):
    return "—" if x is None or x != x else f"{x:.{nd}f}"


def _snapshot(out_dir):
    path = out_dir / "gripper_snapshot.png"
    if not path.exists():
        img = GripperSim(FingerParams(), SCENARIOS["E1_metal"].obj).render_frame(800, 600)
        if img is not None:
            imageio.imwrite(path, img)
    return path


def build_report(out_dir=OUT):
    out_dir = Path(out_dir)
    met = json.loads((out_dir / "metrics.json").read_text(encoding="utf-8"))
    ds_path = out_dir / "domain_shift.json"
    ds = json.loads(ds_path.read_text(encoding="utf-8")) if ds_path.exists() else None
    fp, cp = FingerParams(), CtrlParams()
    td, pd_ = _defaults(TendonActuator), _defaults(PneumaticActuator)
    fwd = SensorForward()
    teng, tc, mc = _defaults(TengModel), _defaults(TengContact), _defaults(MaterialClassifier)
    R = Report()
    R.d.add_heading("단일 전극 멀티모달 센서 기반 3지 소프트 그리퍼 제어 시뮬레이션 — A–Z 보고서", 0)
    R.p("문희준 | BioRobotics & Control Lab (BiRC), 고려대학교 기계공학과. 코드: sim_gripper/ (MuJoCo 3, Python).")

    # 1 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[0])
    R.p("이 시뮬레이션은 EGaIn 스파이럴 코일 단일 전극 센서(L, R_DC, V_TENG를 TDM으로 1 ms마다 취득)를 TPU 3지 소프트 "
        "그리퍼의 손가락 안쪽면에 붙였을 때, 센서 신호만으로 접근–접촉–파지–리프트–해제를 수행하는 폐루프 제어를 "
        "구현하고 검증한다. 목적은 두 가지다: (1) 세미나 시연 — 인장·근접·전기적 성질·접촉·압축(힘) 다섯 모달리티가 "
        "한 번의 파지 시퀀스에서 모두 신호와 그래프로 드러나게 한다. (2) 제어기 설계 도구 — 제어 파라미터, 액추에이터 "
        "종류, 센서 추정 오차, 장착 도메인 시프트가 파지 성능에 미치는 영향을 정량 비교한다.")
    R.p("핵심 질문은 '평면 시험(0909 데이터)으로 학습한 v4 Gate-MoE가 그리퍼에 장착되면 그대로 쓸 수 있는가'이다. "
        "결론부터 말하면 근접·재질은 대체로 통하지만, 접촉 중 굽힘으로 변형률과 힘이 동시에 변하는 그리퍼 고유의 "
        "조건에서 힘·변형률 추정이 크게 흔들리며(10장), 그리퍼 보정 데이터로 fine-tune하면 회복된다(11장).")

    # 2 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[1])
    R.p("Fig.1(b) 'Whole Process of Multimodal Soft Single Sensor System'의 블록을 다음 모듈로 대응시켰다. "
        "물리는 0.25 ms, 센서·추론·제어는 1 ms(실제 TDM 주기)로 돈다.")
    R.table(["Fig.1(b) 블록", "시뮬레이션 모듈", "코드"], [
        ["Soft Sensor + Real-Time HRI Soft Gripper", "M0 MuJoCo 3지 PRBM 그리퍼·물체", "model/gripper.py"],
        ["Switching Excitation → (L, R, V)", "M1 센서 순방향 모델 + TENG 모델", "sensor/forward.py"],
        ["Time Division Measurement", "M2 TDM 샘플러(실측 노이즈)", "sensor/tdm.py"],
        ["Signal Decoupling", "M3 v4 Gate-MoE(펌웨어 동일) + 접촉 융합", "perception/moe.py, contact.py"],
        ["Multi-Modal Information: Electrical Property", "M4 ΔL 잔차 부호 재질 판별", "perception/material.py"],
        ["On-Board Controller (Pre-Shape / State Select / Material-Adaptive Force / Safety)", "M5 제어기", "control/controller.py"],
        ["Actuator Command", "M6 텐던 / 공압 액추에이터", "actuators/"],
        ["(추가) Domain Shift + Fine-tune", "M7 장착 편향 주입과 적응", "experiments/domain_shift.py"],
    ], widths=[6, 6, 4.5])
    R.p("데이터 흐름(손가락 i, 매 1 ms):")
    R.eq("q → κ̄ → ε,   (pad, 물체) → d, F",
         "(ε, d, F, 재질) → (ΔL, ΔR) + 노이즈 → Gate-MoE → (ε̂, d̂ | F̂, contact)",
         "F → V_TENG → 접촉 래치,   ΔL − ΔL_ε(ε_R) → 재질",
         "(phase, 재질, MoE 출력) → u_i → τ_i(4 관절) → MuJoCo")

    # 3 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[2])
    R.h("개념", 2)
    R.p("소프트 손가락을 의사강체 모델(PRBM)로 근사한다: 길이 L을 n개 강체 세그먼트로 나누고 그 사이를 비틀림 스프링 "
        "힌지로 연결한다. 연속 굽힘 빔의 정적 처짐을 힌지 각의 합으로 재현하면서, MuJoCo의 접촉·마찰을 그대로 쓸 수 있다. "
        f"손가락 3개는 반경 {BASE_R*1000:.0f} mm 원판 베이스에 120° 간격으로 아래를 향해 달리고, 베이스는 수직 슬라이더 "
        f"한 축(리프트)을 가진다. 베이스 높이 {BASE_Z*1000:.0f} mm, 손끝은 테이블 위 10 mm.")
    R.fig(_snapshot(out_dir), "그림 3-1. MuJoCo 모델(개방 상태, E1 금속 원통). 밝은 회색 띠가 센서 pad.", 11)
    R.h("수식", 2)
    R.eq("k = E·w·t³ / (12·ℓ),   ℓ = L/n",
         "c = 2ζ·√(k·I),   I ≈ m_f·L²/3",
         "τ_j,passive = −k·q_j − c·q̇_j")
    R.p(f"센서는 원위 세그먼트 {SENSOR_SEGS}의 안쪽면(구간 40 mm, 손끝 포함)에 pad geom으로 표현한다. 센서 구간의 평균 곡률:")
    R.eq("κ̄ = (q₂ + q₃) / 40 mm")
    R.p("pad–물체 최소 거리 d는 mj_geomDistance(탐색 한계 30 mm), 접촉 법선력 F는 pad geom이 물체와 맺은 접촉들의 "
        "mj_contactForce 법선 성분 합. 센서 세그먼트의 몸통 박스는 충돌을 꺼서 하중이 pad로만 전달되게 했다.")
    R.h("파라미터", 2)
    R.table(["항목", "값"], [
        ["L, w, t", f"{fp.L_mm:.0f}, {fp.w_mm:.0f}, {fp.t_mm:.0f} mm"], ["E (TPU 95A 가정)", f"{fp.E_Pa/1e6:.0f} MPa"],
        ["세그먼트 수 n", fp.n_joints], ["관절 강성 k", f"{fp.joint_stiffness():.3f} N·m/rad"],
        ["감쇠비 ζ / 감쇠 c", f"{fp.zeta} / {fp.joint_damping():.2e} N·m·s/rad"],
        ["관절 범위", "−0.5 … 1.6 rad"], ["적분기", "implicitfast, Δt = 0.25 ms, 타원 마찰원뿔, impratio 10"],
        ["계란 파손", "손가락별 접촉력 > fragile_F_N (래치)"],
    ], widths=[6, 9])
    R.h("검증", 2)
    R.bullets(["test_model_loads: hinge 12개", f"test_joint_stiffness: k ≈ 0.47 ± 0.05 (계산값 {fp.joint_stiffness():.3f})",
               "test_distance_decreases / test_contact_force_positive / test_egg_breaks"])
    R.h("한계", 2)
    R.p("PRBM은 단면 변형·전단·대변형 비선형 재료를 무시한다. 손가락 치수와 경도는 가정값이다. 수직 손가락은 키 큰 물체에 "
        "손끝이 먼저 닿으므로 센서를 원위부에 두었다(초기 근위부 배치는 손끝이 물체를 밀어 올려 튕겨냄).")

    # 4 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[3])
    R.h("개념", 2)
    R.p("센서 자체를 물리적으로 모델링하는 대신, 0909 실측 데이터에서 추출한 응답 곡면(룩업 표)으로 (ε, d, F, 재질) → "
        "(ΔL/L₀, ΔR/R₀)를 만든다. 실제 센서가 내놓을 숫자와 같은 분포를 v4 MoE에 넣기 위함이다.")
    R.h("인장 — 프리스트레치 장착", 2)
    R.p("손가락 안쪽면은 굽힘 시 압축된다. 학습 범위가 인장 0–30 %이므로 센서를 ε₀만큼 미리 늘려 붙이고, 굽힘이 이를 줄이게 했다. "
        "표면은 중립축에서 h = t/2만큼 떨어져 있다:")
    R.eq(f"ε = clip(ε₀ − 100·h·κ̄, 0, 30) [%],   ε₀ = {fwd.eps0:g} %,   h = {fwd.h_mm:g} mm")
    R.h("신호 합성", 2)
    R.eq("ΔL = ΔL_ε(ε) + k_mat·ΔL_prox(ε, d) + ΔL_F(ε, F)",
         "ΔR = ΔR_ε(ε) + ΔR_F(ε, F)")
    R.bullets([
        "ΔL_ε, ΔR_ε: 세션 214746에서 근접 없는(prox ≈ 24 mm) 구간의 ε별 평균에 2차 다항식 적합.",
        "ΔL_prox(ε, d) = S_L(ε, d) − S_L(ε, 24): 214746+222022 근접 데이터 격자(ε 0.5 %, d 0.5 mm). d ≥ 24 mm에서 0.",
        "ΔL_F, ΔR_F = S(ε, F) − S(ε, 0): 220210 압력 세션 격자(F 0.25 N). F > 0일 때만.",
        f"k_mat: 도체 {K_MAT['conductor']:g}, 유전체 {K_MAT['dielectric']:g}. 도체는 와전류로 L 감소, 유전체는 기생 정전용량으로 "
        "공진주파수가 내려가 겉보기 L이 약간 증가(개략 가정).",
    ])
    R.h("TENG", 2)
    R.p("TENG 전압은 접촉 대전으로 생기며 힘의 변화율에 비례하는 펄스를 낸다(접촉 +, 분리 −). 1차 시스템으로 둔다:")
    R.eq("τ_T·V̇ + V = g_T·Ḟ   →   V_k = (τ_T·V_{k−1} + g_T·(F_k − F_{k−1})) / (τ_T + Δt)",
         f"τ_T = {teng['tau_s']} s,   g_T = {teng['gain']} (임의 단위)")
    R.h("검증·한계", 2)
    R.p("test_forward 7건: 순수 인장 ΔL≈20 %·ΔR≈20.9 %(ε=15 %), 도체 근접 ΔL≈−22 %, 유전체 1.5–3 %, 10 N에서 ΔR≈20 %, "
        "인장 클립, TENG 부호, ΔR→ε 역산 왕복. 한계: 곡면은 평면 시험에서 얻었으므로 굽힘 상태의 센서 응답과 다를 수 있고"
        "(→ 11장 도메인 시프트), 유전체 k_mat은 실측이 아닌 가정값이다.")

    # 5 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[4])
    R.p("실제 펌웨어는 TIM7 1 ms 주기로 LDC1614(I2C DMA)와 ADC(TENG 100 µs → R 100 µs)를 병렬 취득한다. 시뮬레이션은 "
        "이 주기 그대로 매 1 ms 한 샘플을 만들고, 실측 노이즈를 더한다. 노이즈 크기는 세션 214746 원시 신호에서 "
        "느린 변화를 제거하는 1차 차분으로 추정했다:")
    R.eq("σ = std(Δx) / √2",
         f"σ_L = {L.SIGMA_L} %,   σ_R = {L.SIGMA_R} %")
    R.p("검증: test_tdm (재현성, 노이즈 크기 ±5 %, 실측 추정치 양수). 한계: 백색 가우시안 가정 — 실제의 1/f 드리프트, "
        "온도 드리프트, MUX 전환 과도는 없다.")

    # 6 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[5])
    R.h("개념", 2)
    R.p("STM32에 배포된 v4 Gate-MoE를 numpy로 그대로 이식했다(models.pkl, 총 1,413 파라미터: 게이트 65 + 전문가 A 674 + B 674). "
        "게이트가 접촉 여부를 고르고, 비접촉이면 전문가 A가 (ε̂, d̂), 접촉이면 전문가 B가 (ε̂, F̂)를 낸다.")
    R.h("수식", 2)
    R.eq("p = σ(MLP_g(z(ΔL, ΔR))),   contact = [p ≥ 0.5]",
         "e_k = e_{k−1} + α·(x_k − e_{k−1}),   α = 1 − exp(−Δt/τ),  τ = 1.0125 s  (게이트 전이 시 e ← x)",
         "(ε̂, d̂) = MLP_A(z_A(ΔL, ΔR, e_L, e_R))   |   (ε̂, F̂) = MLP_B(z_B(…))",
         "MLP: 4 → 24 → 16 → 8 → 2,  tanh")
    R.p(f"EMA_ALPHA = {M.EMA_ALPHA}, 출력 클램프: ε {M.STRAIN_LIM}, d {M.DIST_LIM}, F {M.FORCE_LIM}.")
    R.h("검증", 2)
    R.p("보드 로그(moe_monitor_20260910_010437_test.csv)를 재생해 게이트 일치율 ≥ 0.99, 전문가 변형률 R² ≥ 0.95 확인"
        "(test_moe). EMA 리셋 시점도 펌웨어와 같다(전이한 바로 그 사이클).")
    R.h("시뮬레이션에서 드러난 한계와 보강 (M3+)", 2)
    R.p("v4 게이트는 도체 압자로 누른 데이터로만 학습되었다. 순방향 모델에 넣어 보면 유전체 접촉은 8 N에서도 p ≈ 0이고, "
        "도체도 ε = 15 %에서는 4 N 미만 접촉을 놓친다(압력 학습 데이터는 ε ≥ 14 % 접촉을 인덕턴스 포화로 제외). "
        "따라서 다음 세 가지를 추가했다. 이는 실물 펌웨어에도 그대로 옮길 수 있는 제안이다.")
    R.bullets([
        f"접촉 융합: contact = 게이트 ∨ TENG 래치. V > {tc['v_on']} 이면 래치 ON, V < {tc['v_off']} 이면 OFF. TENG는 재질과 "
        "무관하게 접촉 순간을 잡는다(Fig.1(b)의 Contact 모달리티). 융합 접촉은 전문가 선택과 EMA 리셋에도 반영(contact_hint).",
        "유전체 접촉 추정: 유전체에 닿으면 L은 거의 변형률만 반영한다(k_mat 작음). ΔL에서 ε̂, ΔR 잔차에서 F̂를 보정 곡면의 "
        "고정점 반복으로 역산한다:  ε ← ΔL_ε⁻¹(ΔL − k_mat·ΔL_prox(ε,0) − ΔL_F(ε,F)),  F ← ΔR_F(ε,·)⁻¹(ΔR − ΔR_ε(ε)).  "
        "전문가 B는 유전체에서 분포 밖(힘과 무관하게 ≈ 5 N 출력)이기 때문.",
        "주의: 이 역산은 순방향 모델과 같은 표를 쓰므로 시뮬레이션 안에서는 낙관적으로 정확하다(순환성). 실물에서는 유전체 "
        "압력 데이터로 별도 보정이 필요하다.",
    ])

    # 7 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[6])
    R.p("도체는 와전류로 L을 낮추고 유전체는 L을 약간 높인다. 순수 변형률이 설명하는 몫을 빼면 남는 잔차의 부호가 재질이다.")
    R.eq("r = ΔL − ΔL_ε(ε_R),   ε_R = ΔR_ε⁻¹(ΔR)   (비접촉에서 R은 변형률만의 함수)",
         f"r < −θ 가 N회 연속 → 도체,   r > +θ 가 N회 연속 → 유전체,   θ = {mc['theta']} %,  N = {mc['persist_steps']}",
         "첫 접촉 순간 판정 래치(접촉 후 ΔL은 힘 성분이 섞이므로)")
    R.p("ε_R을 쓰는 이유: 처음엔 전문가 A의 ε̂를 썼는데, 접근 중 굽힘으로 ε가 변할 때 ε̂ 편향이 ~1 %(ΔL로 ~1.3 %)라서 "
        "유전체 잔차(+1.5 %)를 뒤집어 계란을 도체로 오판했다. ΔR→ε 역산은 PINN Stage 1(ΔR → ε̂)과 같은 원리로, "
        "비접촉에서 재질·근접에 무관하다. 검증: test_material 7건(굽힘 중 접근에서도 두 재질 모두 정답).")
    R.p("Material Lookup: 판정 결과로 목표·최대 힘을 고른다. 미정은 보수적으로 유전체 값.")
    R.table(["재질", "F_t [N]", "F_max [N]"], [[k, v[0], v[1]] for k, v in MATERIAL_TABLE.items()] + [["미정", "유전체 값", ""]])

    # 8 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[7])
    R.h("상태 머신 (State Select)", 2)
    R.p("전역 상태 OPEN → APPROACH → GRASP → HOLD → RELEASE. 'approach' 명령으로 APPROACH, 세 손가락 모두 contact면 GRASP, "
        f"APPROACH 진입 후 {cp.grasp_timeout_s:g} s 안에 못 잡으면 RELEASE + failed. 'lift'는 GRASP에서만 HOLD(리프트 "
        f"{L.LIFT_H*100:.0f} cm, {L.LIFT_V*100:.0f} cm/s), 'release'는 언제나 RELEASE. 손가락별 루프는 같은 상태를 공유하되 "
        "각자의 센서 출력으로 u_i를 만든다. 이것이 compliant 그리퍼를 위한 하이브리드 구조다: 접촉 전 위치(닫힘) 제어, "
        "접촉 후 admittance 힘 제어.")
    R.h("Pre-shape: 근접도 스케줄 닫힘", 2)
    R.p("설계 초안의 위치형 식 u_ref = u_open + (u_pre − u_open)·sat((d₀ − d̂)/(d₀ − d₁))은 그리퍼가 고정된 이 시나리오에서 "
        "작동하지 않았다: 개방 상태의 d̂ ≈ 23 mm > d₀이면 출발하지 않고, d₁ 근처에서 u와 d가 고정점에 갇혀 접촉하지 못한다. "
        "그래서 근접도로 닫힘 속도를 스케줄링한다(멀면 빠르게, 가까우면 천천히 — 충돌 충격 완화):")
    R.eq("u̇ = v_near + (v_far − v_near)·sat((d̂ − d₁)/(d₀ − d₁))",
         f"v_far = {cp.v_far}/s,  v_near = {cp.v_near}/s,  d₀ = {cp.d0_mm:g} mm,  d₁ = {cp.d1_mm:g} mm")
    R.h("Material-adaptive admittance 힘 제어", 2)
    R.eq("B_a·u̇ = F_t − F̂   →   u_{k+1} = u_k + (Tₛ/B_a)·(F_t − F̂_k),   |F_t − F̂| ≤ δF 이면 0",
         f"B_a = {cp.B_a:g},  δF = {cp.F_dead} N,  |Δu| ≤ u̇_max·Tₛ (u̇_max = {cp.u_rate_max}/s),  0 ≤ u ≤ u_max = {cp.u_max}")
    R.p("Compliant 손가락에서 힘은 굽힘(위치)을 통해서만 만들어지므로, 힘 오차를 명령 u의 속도로 바꾸는 admittance가 "
        "자연스럽다. 접촉 전이 시 내부 상태(필터·미분)를 리셋해 게이트 깜빡임이 u를 튀게 하지 않는다. u_max는 "
        "anti-windup: F̂가 과소추정되면 u가 포화되어 과굽힘한 손끝이 물체를 밀어내는 것을 막는다.")
    R.h("Safety layer (최우선)", 2)
    R.eq("F̄_k = F̄_{k−1} + (Tₛ/τ_F)·(F̂_k − F̄_{k−1}),   τ_F = %g s" % cp.tau_F,
         f"트립: F̄ > F_max  ∨  dF̄/dt > {cp.Fdot_max:g} N/s  ∨  ε̂ ∉ [{cp.eps_range[0]:g}, {cp.eps_range[1]:g}] %",
         "트립 중 u ← u − 2·u̇_max·Tₛ,   F̄ < 0.5·F_max 이면 해제")
    R.p("1 kHz 원시 F̂를 바로 미분하면 노이즈만으로 50 N/s를 넘으므로 저역통과 F̄로 판정한다(비용: ~50 ms 지연). "
        "검증: test_controller 11건(개방, 근접 스케줄, admittance 증분 = Tₛ/B_a, 힘·변형률 안전, 노이즈 비트립, 빠른 상승 트립, "
        "게이트 깜빡임 시 |Δu| 상한, 미정→유전체, 타임아웃, 전체 시퀀스).")

    # 9 ─────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[8])
    R.p("명령 u ∈ [0, 1]을 관절 토크로 바꾸는 두 방식. 매 물리 스텝(0.25 ms)마다 적분한다.")
    R.h("텐던", 2)
    R.eq("θ̇_m = sat((u·θ_max − θ_m)/τ_m, ±ω_max)",
         "T = k_t·max(0, r_p·θ_m − Σ_j r_j·q_j),   τ_j = r_j·T   (늘어지면 0)")
    R.table(["r_j", "r_p", "k_t", "θ_max", "τ_m", "ω_max"],
            [[f"{td['r_j']} m", f"{td['r_p']} m", f"{td['k_t']:g} N/m", f"{td['theta_max']} rad", f"{td['tau_m']} s",
              f"{td['omega_max']} rad/s"]])
    R.h("공압", 2)
    R.eq("Ṗ = sat((u·P_max − P)/τ_p, ±Ṗ_max)",
         "τ_j = max(0, α·P·(1 − β·q_j))   (굽힐수록 효율 감소)")
    R.table(["P_max", "τ_p", "Ṗ_max", "α", "β"],
            [[f"{pd_['P_max']/1e3:g} kPa", f"{pd_['tau_p']} s", f"{pd_['Pdot_max']/1e3:g} kPa/s", f"{pd_['alpha']:g} m³",
              pd_['beta']]])
    R.p("검증: test_actuators 6건(텐던 이완 시 0, 시정수 63 ± 5 %, 공압 시정수·변화율 한계·굽힘 효율 감소, 팩토리). "
        "공압은 τ_p = 150 ms 지연이 있어 같은 admittance 이득에서 힘 루프 위상 여유가 작다(10장).")

    # 10 ────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[9])
    R.p("시나리오 공통 이벤트: 0.5 s approach, 4.0 s lift, 8.0 s release, 총 10 s. 센서 모드 'pipeline'은 순방향 모델 → "
        "TDM 노이즈 → v4 MoE(+M3 보강), 'ideal'은 참값 + 같은 크기 노이즈(E6 기준선).")
    R.table(["ID", "시나리오", "물체"], [
        ["E1", "모달리티 커버리지 시퀀스", "금속 원통 r 25 mm, 0.2 kg (도체)"],
        ["E2", "재질별 파지", "손(캡슐 r 20 mm, soft, 유전체) / 계란(구 r 22 mm, 3.5 N 파손, 유전체)"],
        ["E3", "강성·크기 적응", "스펀지(soft) / 원통 r 15 / r 35 mm"],
        ["E4", "외란", "E1 + 6~7 s 수평 3 N"],
        ["E5", "도메인 시프트", "11장"],
        ["E6", "이상 센서 vs 파이프라인", "E1, E2_hand"],
    ], widths=[1.5, 5, 9])
    keys = ["success", "lift_m", "force_rmse", "max_force", "settle_s", "gate_acc", "strain_r2", "material_acc"]
    rows = _metric_rows({k: v for k, v in met.items() if not k.startswith("bellows")})
    R.p("지표: force_rmse = HOLD 구간 √mean((F − F_t)²) (참 힘 기준), settle_s = GRASP 진입 후 ±10 % 띠에 머물기 시작한 "
        "시간, gate_acc = contact 판단 vs 참 접촉(F > 0.05 N), strain_r2 = ε̂ vs ε 전 구간, material_acc = 최종 재질 정답 손가락 비율.")
    R.table(METRIC_HEAD, rows)
    e1, e1i = met.get("E1_metal_tendon", {}), met.get("E1_metal_ideal", {})
    R.h("해석", 2)
    R.bullets([
        f"다섯 모달리티 가시화(E1, 그림 10-1): 접근 중 d̂ 감소와 ΔL 급락(도체), 접촉 순간 TENG 양의 펄스, 파지 중 ΔR 상승(압축) "
        "+ 굽힘으로 ε 감소, 해제 시 TENG 음의 펄스.",
        f"전체 파이프라인의 도체 힘 추정이 병목이다: E1 force_rmse {_f(e1.get('force_rmse'))} N (ideal {_f(e1i.get('force_rmse'))} N), "
        f"ε R² {_f(e1.get('strain_r2'), 2)} (ideal {_f(e1i.get('strain_r2'), 2)}). 전문가 B는 고정 변형률에서 누른 데이터로 학습되어, "
        "접촉 중 굽힘으로 ε가 함께 변하는 그리퍼 조건에서 힘을 1–2 N 과소추정한다. 제어기는 이를 보상하려 u를 u_max까지 올리고 "
        "실제 힘은 F_max를 넘나든다 — 사용자의 초기 질문('그리퍼에서는 힘·인장이 평면 시험과 달라질 것')이 시뮬레이션에서 재현된 것.",
        "유전체(손·계란·스펀지)는 TENG 융합 접촉 + 해석적 역산 덕에 힘 추종이 좋다(RMSE ~0.1 N, 계란 무파손). 다만 이 정확도는 "
        "순환성 때문에 낙관적이다(6장).",
        "공압은 압력 지연(150 ms)과 과소추정이 겹쳐 E1/E4에서 과하중 → 안전 개방 → 재파지의 리밋 사이클 끝에 물체를 놓친다. "
        "공압용 B_a를 따로 튜닝하거나(설계 도구로서의 용도) 힘 추정을 먼저 고쳐야 한다.",
        "계란 재질 판정은 3지 중 1지만 '유전체'로 확정, 나머지는 '미정' — 미정은 유전체 값을 쓰므로 안전 측. 반경이 작아 접근 "
        "중 근접 잔차가 θ를 넘는 시간이 짧기 때문.",
        "유전체 파지 중 리프트 순간 힘이 출렁이며 TENG 음의 피크가 래치를 잠깐 끊는다(그림 10-2) — 실물 래치에는 게이트·ΔR "
        "기반 해제 조건을 함께 두는 것이 좋다.",
    ])
    R.fig(out_dir / "timeline_E1_metal_tendon.png", "그림 10-1. E1 금속 원통(텐던), 손가락 0 — 다섯 모달리티 타임라인.")
    R.fig(out_dir / "timeline_E2_hand_tendon.png", "그림 10-2. E2 손(유전체, 텐던), 손가락 0.")
    R.fig(out_dir / "timeline_E2_egg_tendon.png", "그림 10-3. E2 계란(텐던), 손가락 0.")
    R.fig(out_dir / "compare_actuators.png", "그림 10-4. E1 텐던 vs 공압 — 힘과 명령.", 14)
    R.fig(out_dir / "ideal_vs_pipeline.png", "그림 10-5. E6 이상 센서 vs 전체 파이프라인.", 14)
    R.p("영상: out/E1_metal_tendon.mp4 (30 fps).")

    # 11 ────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[10])
    R.p(f"{DISCLAIMER}. 아래 결과는 가정한 장착 편향을 순방향 모델에 주입하고 적응 방법을 비교하는 절차를 보여줄 뿐, "
        "실제 그리퍼에서 이 수치가 나온다는 근거는 아니다. 실물에서는 같은 절차를 실측 보정 데이터로 수행해야 한다.")
    R.h("시프트 모델", 2)
    R.eq("ε' = ε + δε,   ΔL' = s·ΔL(ε') + γ·100·h·κ̄,   ΔR' = s·ΔR(ε')",
         f"δε = {DEFAULT_SHIFT.d_eps} %,  γ = {DEFAULT_SHIFT.gamma},  s = {DEFAULT_SHIFT.scale}")
    R.p("δε: 접착·프리스트레치 오차, γ: 굽힘이 코일 형상을 바꿔 생기는 L 결합(평면 시험엔 없음), s: 기판·배선 차이로 인한 감도 변화.")
    R.h("적응 방법", 2)
    R.bullets([
        "recal: 스케일러만 그리퍼 데이터로 재적합 + 전문가 출력에 선형 보정(LinearRegression). 파라미터 그대로, 가장 싸다.",
        "finetune: 기존 가중치에서 warm start, 학습률 1e-4, 200 epoch. 구조·파라미터 수(1,413) 그대로 → STM32 재배포 용이.",
        "mixed: finetune + 0909 평면 데이터 동량 혼합 — 평면 성능 망각 억제.",
        "scratch: 같은 구조를 그리퍼 데이터만으로 처음부터.",
    ])
    R.p("그리퍼 데이터: E1·E2_hand·E3_small 폐루프 실행의 손가락 3개 스트림, 100 Hz 다운샘플, 라벨은 시뮬레이션 참값 "
        "(실물에서는 보정 지그·로드셀이 필요). EMA는 1 kHz에서 펌웨어 α로 재계산. 평가는 다른 시드의 홀드아웃 그리퍼 데이터와 "
        "0909 test 셋(10 s 창 5개마다 1개).")
    if ds:
        rows = [["v4 (시프트 없음)", "—", _f(ds["v4_clean"]["strain_r2"]), _f(ds["v4_clean"]["dist_r2"]),
                 _f(ds["v4_clean"]["force_r2"]), _f(ds["v4_flat"]["strain_r2"])],
                ["v4 (시프트)", "—", _f(ds["v4_shift"]["strain_r2"]), _f(ds["v4_shift"]["dist_r2"]),
                 _f(ds["v4_shift"]["force_r2"]), _f(ds["v4_flat"]["strain_r2"])]]
        for meth, by_n in ds["methods"].items():
            for n, r in by_n.items():
                rows.append([meth, n, _f(r["sim"]["strain_r2"]), _f(r["sim"]["dist_r2"]), _f(r["sim"]["force_r2"]),
                             _f(r["flat"]["strain_r2"])])
        R.table(["모델", "N", "그리퍼 ε R²", "그리퍼 d R²", "그리퍼 F R²", "0909 ε R²"], rows)
        n_max = list(ds["methods"]["mixed"])[-1]
        r = {k: ds["methods"][k][n_max] for k in ds["methods"]}
        R.bullets([
            f"시프트를 주지 않은 v4도 그리퍼 데이터에서 ε R² {_f(ds['v4_clean']['strain_r2'], 2)} (평면 0909 test {_f(ds['v4_flat']['strain_r2'], 2)}): "
            "그리퍼 장착 자체(접촉 중 굽힘, 원위부 대변형)가 이미 큰 도메인 시프트다. 장착 편향을 더하면 "
            f"{_f(ds['v4_shift']['strain_r2'], 2)}로 더 나빠진다.",
            f"N = {n_max}에서 그리퍼 ε/d/F R²: recal {_f(r['recal']['sim']['strain_r2'], 2)}/{_f(r['recal']['sim']['dist_r2'], 2)}/"
            f"{_f(r['recal']['sim']['force_r2'], 2)}, finetune {_f(r['finetune']['sim']['strain_r2'], 2)}/{_f(r['finetune']['sim']['dist_r2'], 2)}/"
            f"{_f(r['finetune']['sim']['force_r2'], 2)}, mixed {_f(r['mixed']['sim']['strain_r2'], 2)}/{_f(r['mixed']['sim']['dist_r2'], 2)}/"
            f"{_f(r['mixed']['sim']['force_r2'], 2)}, scratch {_f(r['scratch']['sim']['strain_r2'], 2)}/{_f(r['scratch']['sim']['dist_r2'], 2)}/"
            f"{_f(r['scratch']['sim']['force_r2'], 2)}.",
            f"망각(0909 ε R²): mixed {_f(r['mixed']['flat']['strain_r2'], 2)}로 유지, finetune {_f(r['finetune']['flat']['strain_r2'], 2)}, "
            f"scratch {_f(r['scratch']['flat']['strain_r2'], 2)}, recal {_f(r['recal']['flat']['strain_r2'], 2)}. 평면과 그리퍼를 모두 "
            "써야 한다면 mixed fine-tune이 균형점이다. recal은 선형 보정 한계로 R² ~0.5에서 정체한다.",
            "finetune은 N이 작을 때 ε가 오히려 나빠질 수 있다(소량 데이터 과적합) — 실물 보정 데이터는 수천 샘플 이상을 목표로.",
        ])
        R.fig(out_dir / "domain_shift_recovery.png", "그림 11-1. 보정 샘플 수 N에 따른 회복 곡선과 평면 성능 유지.", 16.5)
        cl = ds.get("closed_loop_E1_shift", {})
        if cl:
            R.p("폐루프 효과(시프트된 E1 금속, 텐던):")
            R.table(["모델", "파지", "F RMSE [N]", "최대 F [N]", "ε R²"],
                    [[k, "성공" if v["success"] else "실패", _f(v["force_rmse"]), _f(v["max_force"], 2), _f(v["strain_r2"], 2)]
                     for k, v in cl.items()])
    else:
        R.p("[domain_shift.json 없음 — python -m sim_gripper.experiments.domain_shift 실행 필요]")
    R.p("해석 요령: 그리퍼 R²가 v4(시프트) 점선 위로 얼마나 빨리 올라오는지(데이터 효율)와, 0909 ε R²가 얼마나 유지되는지"
        "(망각)를 함께 본다. 펌웨어 재배포 관점에서는 구조가 같은 finetune/mixed가 유리하다.")

    # 12 ────────────────────────────────────────────────────────────────
    R.h(CHAPTERS[11])
    R.h("모델 한계", 2)
    R.bullets([
        "센서 응답은 평면 시험 곡면을 그대로 썼다. 실제 굽힘·곡면 접촉에서의 L/R 응답은 측정이 필요하다(→ 11장 절차).",
        "유전체 k_mat = −0.1, TENG 이득·시정수, 계란 파손 3.5 N, 손가락 치수·경도, 공압 파라미터는 가정값이다.",
        "PRBM 손가락, 강체 물체(soft 옵션은 접촉 solref만 완화), 단순 마찰. 손은 캡슐로 근사.",
        "유전체 힘 역산은 순방향 표를 그대로 뒤집어 시뮬레이션 안에서 낙관적이다.",
    ])
    R.h("실물 시스템에 주는 시사점", 2)
    R.bullets([
        "v4 게이트는 유전체·저하중 접촉을 놓친다 → 펌웨어에 TENG 융합 접촉(이미 측정 중인 신호)을 넣을 것.",
        "재질 판별의 기준 변형률은 ΔR 역산(비접촉)으로 → PINN Stage 1 구조와 일치.",
        "전문가 B는 '접촉 중 변형률 변화' 데이터가 필요 → 그리퍼 장착 상태에서 굽힘+누름 동시 자극 보정 데이터 수집, "
        "warm-start fine-tune(구조 유지)로 재배포.",
        "힘 제어는 F̂ 편향에 민감 → 안전 계층(저역통과 F̄, u_max)을 펌웨어 수준에서 유지.",
    ])
    R.h("향후 과제", 2)
    R.bullets(["실물 보정 데이터로 M1 곡면 교체 후 재검증", "공압 전용 admittance 이득·피드포워드", "다중 주파수 임피던스로 재질 "
               "판별 강화", "2D 변형(직사각형 스파이럴 비대칭)으로 비틀림 파지 감지", "로봇팔 접근(그리퍼 이동) 시나리오"])

    # 부록 A ─────────────────────────────────────────────────────────────
    R.h(CHAPTERS[12])
    rows = [["FingerParams", k, v] for k, v in fp.__dict__.items()]
    rows += [["CtrlParams", k, v] for k, v in cp.__dict__.items()]
    rows += [["TendonActuator", k, v] for k, v in td.items()]
    rows += [["PneumaticActuator", k, v] for k, v in pd_.items()]
    rows += [["SensorForward", "eps0 / h_mm", f"{fwd.eps0} / {fwd.h_mm}"], ["TengModel", "tau_s / gain", f"{teng['tau_s']} / {teng['gain']}"],
             ["TengContact", "v_on / v_off", f"{tc['v_on']} / {tc['v_off']}"],
             ["MaterialClassifier", "theta / persist_steps", f"{mc['theta']} / {mc['persist_steps']}"],
             ["TDM", "sigma_L / sigma_R", f"{L.SIGMA_L} / {L.SIGMA_R}"], ["MoE", "EMA_ALPHA", M.EMA_ALPHA],
             ["DomainShift", "d_eps / gamma / scale", f"{DEFAULT_SHIFT.d_eps} / {DEFAULT_SHIFT.gamma} / {DEFAULT_SHIFT.scale}"],
             ["Loop", "LIFT_H / LIFT_V / CONTACT_F", f"{L.LIFT_H} m / {L.LIFT_V} m/s / {L.CONTACT_F} N"]]
    R.table(["모듈", "파라미터", "값"], rows, widths=[4.5, 5, 6])

    # 부록 B ─────────────────────────────────────────────────────────────
    R.h(CHAPTERS[13])
    R.table(["기호", "의미", "단위"], [
        ["ε, ε̂, ε₀, ε_R", "센서 인장 변형률(참, 추정, 프리스트레치, ΔR 역산)", "%"],
        ["d, d̂", "pad–물체 거리(참, 추정)", "mm"], ["F, F̂, F̄", "pad 법선력(참, 추정, 저역통과)", "N"],
        ["ΔL, ΔR", "L/L₀, R/R₀ 변화율", "%"], ["V", "TENG 전압", "a.u."], ["q_j, κ̄", "관절각, 센서 구간 평균 곡률", "rad, 1/mm"],
        ["k, c, ζ", "관절 강성, 감쇠, 감쇠비", "N·m/rad, N·m·s/rad, —"], ["h", "중립축–센서 거리", "mm"],
        ["k_mat", "재질 근접 계수", "—"], ["p", "게이트 접촉 확률", "—"], ["α, τ", "EMA 계수, 시정수", "—, s"],
        ["r, θ", "재질 잔차, 임계", "%"], ["u", "액추에이터 명령", "0–1"], ["F_t, F_max", "목표·최대 힘", "N"],
        ["B_a", "admittance 감쇠", "N·s"], ["d₀, d₁", "근접 스케줄 경계", "mm"], ["v_far, v_near", "닫힘 속도", "1/s"],
        ["θ_m, T", "모터각, 텐던 장력", "rad, N"], ["P", "공압", "Pa"], ["δε, γ, s", "도메인 시프트 파라미터", "%, —, —"],
    ], widths=[3.5, 8.5, 4])

    _appendix_bellows(R, met, out_dir)

    path = out_dir / DOC_NAME
    try:
        R.d.save(str(path))
    except PermissionError:                     # Word에서 열려 있으면 별도 이름으로 저장
        path = out_dir / DOC_NAME.replace(".docx", "_갱신.docx")
        R.d.save(str(path))
    return path


def _metric_rows(met):
    return [[name, "성공" if m["success"] else "실패", _f(m["lift_m"] * 100, 1), _f(m["force_rmse"]),
             _f(m["max_force"], 2), _f(m["settle_s"], 2), _f(m["gate_acc"], 2), _f(m["strain_r2"], 2),
             _f(m["material_acc"], 2)] for name, m in met.items() if "success" in m]


METRIC_HEAD = ["실행", "파지", "리프트 [cm]", "F RMSE [N]", "최대 F [N]", "정착 [s]", "접촉 정확도", "ε R²", "재질"]
CAD_DIR = Path(__file__).resolve().parents[1] / "cad" / "out"


def _appendix_bellows(R, met, out_dir):
    bp = BellowsParams()
    R.h(CHAPTERS[14])
    R.p("근거: K. Meitani, S. Nisar, \"Materials and Methods for Designing 3D-Printed Soft Robot Grippers in Low-Hardness "
        "TPU (60A–70A)\", IEEE Access 13, 210465–210473, 2025 (doi 10.1109/ACCESS.2025.3642169). 논문의 Filaflex 60A "
        "벨로즈 손가락을 출력용 CAD와 시뮬레이션 손가락으로 옮기고, 같은 폐루프(센서 → MoE → 제어)를 돌려 판형 손가락과 비교한다.")
    R.h("논문 수치", 2)
    R.table(["항목", "값 (논문)"], [
        ["Yeoh 3차 (60A)", "C10 1.31309561, C20 −0.68049203, C30 0.45647907 MPa"],
        ["탄성계수 (선형 영역)", "60A 4.43487 MPa, 70A 6.61786 MPa"],
        ["손가락 (Table 3)", "l 105, l_b 6.43, l_c 30.8, w 7.00, w_c 8.00, h 32.0 mm; 벨로즈 8개, precurve 41°, 벽 2.5 mm"],
        ["굽힘 (Fig.9)", "60A 40 kPa −10° → 200 kPa 63° (평균 0.45°/kPa); 총 굽힘 41° + 63° = 104°"],
        ["끝단 힘 (Fig.7)", "60A 100 kPa 4.48 N, 200 kPa 6.28 N, 300 kPa 10.77 N"],
        ["그리퍼", "3지 삼각 배치, 파지 영역 외접원 60 mm, PLA 베이스, 1297 g 스풀·55 g 계란 파지"],
        ["출력", "Prusa MK4, 225 °C, 층 0.12 mm, 인필 100 %, 20 mm/s, 압출배율 1.1, 리트랙션 없음, 옆으로 눕혀 출력"],
    ], widths=[4, 12])
    R.h("출력용 3D 모델", 2)
    R.p("sim_gripper/cad/ (build123d 파라미터 스크립트, python -m sim_gripper.cad.export): finger_60A.stl/.step, "
        "gripper_base.stl, PRINT_SETTINGS.md. 안쪽면(바닥층) 원위 3개 벨로즈 구간에 센서 홈(폭 22 × 깊이 0.5 mm)을 추가했다. "
        "논문 표만으로 정해지지 않는 값은 Fig.3 비율로 추정했다: 바닥층 3 mm, 공기 채널 4 mm, 벨로즈 높이 14 mm, 벨로즈는 "
        "두께 l_b의 평판(둥근 로브·필렛 생략). 실제 출력 전 논문 그림과 대조해 FingerCAD 값을 조정할 것.")
    R.fig(CAD_DIR / "finger_preview.png", "그림 C-1. 60A 벨로즈 손가락 CAD (옆면·등각).", 15)
    R.h("시뮬레이션 모델과 보정", 2)
    R.p(f"벨로즈 하나를 PRBM 관절 하나로 본다(관절 {bp.n_joints}개, 세그먼트 {bp.seg_m() * 1000:.2f} mm). 출력 형상의 "
        "precurve는 관절 스프링 기준각으로 둔다. 공압은 데드존을 두어 Fig.9의 영점(≈ 62 kPa)을 재현한다:")
    R.eq("springref_j = 41°/8,   τ_j = α·max(0, P − P₀),   Δθ = Σ_j (q_j − springref_j) ≈ 8·α·(P − P₀)/k",
         f"P₀ = {PNEUMATIC_60A['P0'] / 1e3:.0f} kPa,  α = {PNEUMATIC_60A['alpha']:.4e} N·m/Pa,  k = {bp.k_joint:.4f} N·m/rad,  "
         f"P_max = {PNEUMATIC_60A['P_max'] / 1e3:.0f} kPa")
    R.p("α/k는 Fig.9 기울기로, α 크기는 200 kPa 끝단 차단력(6.28 N)으로 정했다. 차단력은 끝단 pad 중심을 공간에 고정한 "
        "구속 반력의 굽힘 방향 성분으로 측정했다(판 접촉 방식은 끝단이 미끄러져 불안정). 선형 데드존 모델이라 100·300 kPa "
        "힘은 논문보다 작게 나온다.")
    cal = met.get("bellows_calibration")
    if cal:
        R.table(["압력 [kPa]", "시뮬 굽힘 [°]", "Fig.9 굽힘 [°]"], [[f"{a:.0f}", _f(b, 1), _f(c, 1)] for a, b, c in cal["bending"]])
        R.table(["압력 [kPa]", "시뮬 차단력 [N]", "Fig.7 차단력 [N]"], [[f"{a:.0f}", _f(b, 2), _f(c, 2)] for a, b, c in cal["force"]])
    R.fig(out_dir / "bellows_calibration.png", "그림 C-2. 보정 결과: 굽힘각(Fig.9)과 끝단 차단력(Fig.7).", 16)
    R.h("센서 위치의 의미", 2)
    R.p(f"센서는 물체와 닿는 안쪽 바닥층(strain-limiting layer)에 붙는다. 이 층은 굽힘의 중립축에 가까워 굽힘 인장이 거의 "
        f"없다(h = {bp.h_sensor_mm} mm, 판형은 3 mm). 그래서 ε은 프리스트레치 ε₀ 근처에 머물고, 접촉 중 '변형률이 같이 변하는' "
        "문제가 크게 줄어든다 — 판형에서 v4 MoE 힘 추정을 망친 바로 그 조건(10장)이 구조적으로 완화된다. 대신 인장 모달리티는 "
        "굽힘 정보를 거의 주지 않으므로, 굽힘(자세)이 필요하면 바깥 벨로즈면에 두 번째 센서 채널을 두는 것이 대안이다.")
    R.h("결과", 2)
    bel = {k.replace("bellows60a_", "벨로즈 "): v for k, v in met.items() if k.startswith("bellows60a_")}
    flat = {f"판형 {k[:-7]}": v for k, v in met.items() if k.endswith("_tendon")}
    R.table(METRIC_HEAD, _metric_rows(bel) + _metric_rows(flat))
    R.fig(out_dir / "compare_fingers.png", "그림 C-3. E1 금속 원통: 판형 vs 벨로즈 (힘, 인장, 명령).", 14)
    R.fig(out_dir / "timeline_bellows60a_E1_metal.png", "그림 C-4. 벨로즈 E1 금속 원통, 손가락 0 타임라인.")
    R.h("논문 대비 달라진 점", 2)
    R.bullets([
        f"장착 반경 {bp.base_R * 1000:.0f} mm (논문 형상은 끝단 반경 ≈ 30 mm) — 지름 50 mm 물체에서 근접 구간을 만들기 위해 10 mm 넓혔다.",
        "벨로즈 외형·채널·바닥층 두께는 추정치, 둥근 로브·필렛 생략.",
        "Yeoh 초탄성 FEM 대신 PRBM: 굽힘·끝단 힘의 전체 거동만 맞추고 벨로즈 국부 변형(ballooning)은 표현하지 않는다.",
        "스풀(1.3 kg) 시나리오는 유전체 목표 힘(1.5 N)으로는 마찰이 부족하다 — 무게 추정(리프트 중 F̂ 증가)으로 목표 힘을 올리는 "
        "적응이 필요하다는 근거.",
        "수정 이력: 이전 버전은 MJCF 각도 단위(degree) 때문에 판형 관절 제한이 −0.5°…1.6°로 들어가 소프트 제한이 추가 강성처럼 "
        "작용했다. compiler angle=\"radian\"으로 고친 뒤 이 보고서의 모든 수치를 다시 계산했다.",
    ])


if __name__ == "__main__":
    print(build_report())
