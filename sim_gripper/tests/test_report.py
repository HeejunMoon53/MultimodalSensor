import json

import pytest
from docx import Document

from sim_gripper.experiments.run_all import OUT
from sim_gripper.report.build_report import CHAPTERS, DISCLAIMER, build_report


@pytest.fixture(scope="module")
def doc():
    if not (OUT / "metrics.json").exists():
        pytest.skip("run_all 먼저 실행 필요")
    return Document(str(build_report(OUT)))


def test_report_has_all_chapters(doc):
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(CHAPTERS) == 14 and all(c in heads for c in CHAPTERS)


def test_report_has_disclaimer(doc):
    ps = doc.paragraphs
    k = next(i for i, p in enumerate(ps) if p.text == CHAPTERS[10])
    first = next(p.text for p in ps[k + 1:] if p.text.strip())
    assert DISCLAIMER in first


def test_report_numbers_from_metrics(doc):
    m = json.loads((OUT / "metrics.json").read_text(encoding="utf-8"))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert f"{m['E1_metal_tendon']['force_rmse']:.3f}" in text
