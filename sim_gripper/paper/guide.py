"""논문 Figure 구성 가이드 (docx). 수치는 그림 생성 산출물(CSV·npz·json)에서 읽는다.

python -m sim_gripper.paper.guide  →  sim_gripper/out/paper_figs/논문_Figure_구성가이드.docx
"""
import json
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from sim_gripper.experiments.run_all import OUT
from sim_gripper.paper.diagrams import BUDGET, INFER_US

FIG_DIR = OUT / "paper_figs"
GUIDE = "논문_Figure_구성가이드.docx"
FIGS = [  # (key, 제목)
    ("fig1_architecture", "Fig. 1  전체 알고리즘 구조"),
    ("fig2_timing_budget", "Fig. 2  1 ms 주기 타이밍과 모델 예산"),
    ("fig3_state_machine", "Fig. 3  상태 머신과 모달리티의 역할"),
    ("fig4_grasp_timeline", "Fig. 4  한 번의 파지에서 다섯 모달리티"),
    ("fig5_material_adaptive", "Fig. 5  재질 인식 기반 힘 목표 자동 선택"),
    ("fig6_latency", "Fig. 6  온보드 폐루프의 지연 이점"),
    ("fig7_analytic_vs_moe", "Fig. 7  해석적 역산 vs Gate-MoE"),
    ("fig8_ablation", "Fig. 8  모달리티 ablation"),
    ("fig9_adaptation", "Fig. 9  장착 도메인 적응"),
]


class Doc:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = "맑은 고딕", Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        for s in ("Heading 1", "Heading 2", "Heading 3", "Title"):
            self.d.styles[s].font.name = "맑은 고딕"
            self.d.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            self.d.styles[s].font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        for sec in self.d.sections:
            sec.left_margin = sec.right_margin = Cm(2.2)

    def h(self, t, lv=1):
        self.d.add_heading(t, level=lv)

    def p(self, t, lead=None, italic=False):
        para = self.d.add_paragraph()
        if lead:
            para.add_run(lead).bold = True
        r = para.add_run(t)
        r.italic = italic
        return para

    def bl(self, items):
        for it in items:
            self.d.add_paragraph(it, style="List Bullet")

    def table(self, head, rows):
        t = self.d.add_table(rows=1, cols=len(head))
        t.style = "Light Grid Accent 1"
        for c, x in zip(t.rows[0].cells, head):
            c.text = str(x)
            c.paragraphs[0].runs[0].bold = True
        for r in rows:
            for c, x in zip(t.add_row().cells, r):
                c.text = str(x)
        self.d.add_paragraph()

    def fig(self, key, width=16):
        path = FIG_DIR / f"{key}.png"
        if path.exists():
            self.d.add_picture(str(path), width=Cm(width))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            self.p(f"[{key}.png 없음 — python -m sim_gripper.paper.figures 먼저 실행]", italic=True)


def _numbers():
    n = {}
    lat = pd.read_csv(FIG_DIR / "fig6_latency_data.csv") if (FIG_DIR / "fig6_latency_data.csv").exists() else None
    if lat is not None:
        flat = lat[lat.finger == "flat"]
        n["lat_safe"] = {B: int(g[~g.broken].latency.max()) if (~g.broken).any() else None
                         for B, g in flat.groupby("B_a")}
        n["lat_b1"] = lat[(lat.finger == "flat") & (lat.B_a == 1.0)].set_index("latency").peak.to_dict()
        n["lat_bel"] = lat[lat.finger == "bellows60a"].set_index("latency").peak.to_dict()
        n["lat_broken_b1"] = lat[(lat.finger == "flat") & (lat.B_a == 1.0) & lat.broken].latency.tolist()
    abl = pd.read_csv(FIG_DIR / "fig8_ablation_data.csv") if (FIG_DIR / "fig8_ablation_data.csv").exists() else None
    n["abl"] = abl
    if (FIG_DIR / "fig7_maps.npz").exists():
        m = np.load(FIG_DIR / "fig7_maps.npz")
        near = m["d"] < 4
        n["an_fail_near"] = float(m["fail_rate"][:, near].mean())
        n["an_fail_far"] = float(m["fail_rate"][:, ~near].mean())
        n["moe_err_near"] = float(m["moe_d_err"][:, near].mean())
        n["an_err_mid"] = float(np.nanmean(m["d_err"][:, ~near]))
        n["moe_err_mid"] = float(m["moe_d_err"][:, ~near].mean())
    ds = OUT / "domain_shift.json"
    n["ds"] = json.loads(ds.read_text(encoding="utf-8")) if ds.exists() else None
    return n


def _abl_row(abl, cond, scen):
    r = abl[(abl.condition == cond) & (abl.scenario == scen)].iloc[0]
    return r


def build_guide(out_dir=FIG_DIR):
    n = _numbers()
    D = Doc()
    D.d.add_heading("논문 Figure 구성 가이드 — 단일 전극 멀티모달 센서 기반 Embedded-AI 소프트 그리퍼 제어", 0)
    D.p("모든 그림은 sim_gripper/out/paper_figs/에 PNG(300 dpi)와 SVG로 있다. 라벨은 영문(투고용), 이 문서는 "
        "각 그림의 메시지·패널·데이터 출처·캡션 초안·본문 초안·실물 실험으로 바꿀 때 필요한 것을 정리한다. "
        "재생성: python -m sim_gripper.paper.figures → python -m sim_gripper.paper.guide.")

    D.h("1. 논문이 보여줘야 할 세 가지 주장")
    D.bl([
        "C1 — 단일 전극 하나로 얻는 다섯 모달리티(인장·근접·힘·접촉·재질)가 제어의 모든 결정에 실제로 쓰인다 "
        "(근접 → 접근 속도, 재질 → 목표 힘, 접촉 → 상태 전이, 힘 → admittance, 인장/힘 → 안전). → Fig. 3, 4, 5, 8",
        f"C2 — 이 디커플링은 MCU 위에서 돈다: 약 1.4 k 파라미터, 5.5 KB, 실측 {INFER_US:.0f} µs로 1 ms TDM 주기 안. "
        "해석적 역산이 발산하는 근거리에서도 MoE는 안정적이다. → Fig. 1, 2, 7",
        "C3 — 측정·추론·제어가 같은 칩 위 1 kHz로 닫혀서, PC 경유 지연에서는 불가능한 빠른(높은 이득) 힘 제어를 안전하게 쓸 수 있다. "
        "→ Fig. 6",
    ])
    D.p("Fig. 9(도메인 적응)는 '그리퍼 장착 시 생기는 분포 변화를 같은 구조의 fine-tune으로 회복'을 보여주는 논의/보충 그림이다.")

    D.h("2. 권장 논문 구성과 그림 배치")
    D.table(["절", "내용", "그림/표"], [
        ["1 Introduction", "다중 모달 촉각 필요성, 신호 결합 문제, 장비 의존", "기존 Fig.1(b) 개요"],
        ["2 Sensor & TDM", "EGaIn 코일, 스위칭 여기, L/R/V 시분할", "Fig. 2 (상단 타이밍)"],
        ["3 Embedded decoupling", "Gate-MoE 구조, EMA, TENG 융합, 재질 판별", "Fig. 1, Fig. 7, Table 1"],
        ["4 Control architecture", "상태 머신, 근접 스케줄 pre-shape, 재질 적응 admittance, 안전 계층", "Fig. 3"],
        ["5 Experiments", "파지 시퀀스, 재질 적응, ablation, 지연", "Fig. 4, 5, 8, 6, Table 2"],
        ["6 Discussion", "장착 도메인 적응, 한계, 하드웨어(벨로즈 60A) 변형", "Fig. 9"],
    ])

    D.h("3. Figure별 상세")
    abl, ds = n.get("abl"), n.get("ds")
    for key, title in FIGS:
        D.h(title, 2)
        D.fig(key, 16)
        spec = _SPEC[key](n)
        D.p(spec["msg"], lead="메시지: ")
        D.p(spec["panels"], lead="패널: ")
        D.p(spec["data"], lead="데이터 출처: ")
        D.p(spec["caption"], lead="Caption (EN): ", italic=True)
        D.p(spec["body"], lead="본문 초안: ")
        D.p(spec["real"], lead="실물 실험으로 바꿀 때: ")

    D.h("4. 표")
    D.h("Table 1  On-board 모델 예산 (X-CUBE-AI 리포트 + 보드 실측)", 2)
    D.table(["network", "params", "MACC", "Flash [B]", "RAM [B]"],
            [[k, v["params"], f"{v['macc']:,}", f"{v['flash']:,}", v["ram"]] for k, v in BUDGET.items()] +
            [["total / per inference", sum(v["params"] for v in BUDGET.values()), "1,356",
              f"{sum(v['flash'] for v in BUDGET.values()):,}", 160]])
    D.p(f"추론 지연: 중앙값 {INFER_US:.0f} µs, 최대 265 µs (STM32G473 @170 MHz, float32, 보드 로그 latency_us). "
        "sklearn 기준 파라미터 수는 1,413(게이트 65)이고 X-CUBE-AI는 후처리 3개를 포함해 1,416으로 센다 — 논문에는 한쪽으로 통일.")
    if abl is not None:
        D.h("Table 2  모달리티 ablation (판형 텐던, 그리퍼 보정 MoE, 시뮬레이션)", 2)
        rows = []
        for cond in abl.condition.unique():
            r = [cond]
            for s in ("E2_egg", "E2_hand", "E1_metal"):
                x = _abl_row(abl, cond, s)
                r.append(f"{x.peak:.2f} N {'✓' if x.success else ('✕ broken' if x.broken else '✕')}")
            rows.append(r)
        D.table(["condition", "egg (≤3.5 N)", "human finger (≤3 N)", "metal (≤8 N)"], rows)
    D.h("Table 3  선행 연구 비교 (틀 — 문헌값은 직접 채움)", 2)
    D.table(["연구", "모달리티", "전극 수", "계측 장비", "디커플링", "온보드 실시간 제어"], [
        ["This work", "strain, proximity, force, contact, material", "1 (2 wires)", "custom PCB (LDC1614 + ADC)",
         "Gate-MoE (1.4 k params)", "yes, 1 kHz"],
        ["[ref]", "", "", "LCR meter / …", "", ""],
    ])

    D.h("5. 시뮬레이션 결과를 논문에 쓰는 방법")
    D.bl([
        "Fig. 4–6, 8, 9의 수치는 시뮬레이션이다. 캡션에 'simulated'를 명시하고, 실측 그림이 생기면 같은 레이아웃으로 교체한다 "
        "(그림 코드가 로그 DataFrame만 읽도록 되어 있어 실측 CSV를 같은 열 이름으로 넣으면 재사용 가능).",
        "센서 신호는 0909 실측 곡면 + 실측 노이즈로 합성했다. 추정기는 v4와 같은 구조를 시뮬 그리퍼 데이터로 보정한 모델 — "
        "논문에는 '그리퍼 장착 보정 후'로 쓰고, 보정 전(v4) 결과는 Fig. 9/보충자료에 둔다.",
        "Fig. 1·2·7과 Table 1은 실측/실제 모델 기반이라 그대로 본문에 쓸 수 있다 (Fig. 7의 입력은 0909 실측 곡면 기반 합성).",
        "판형 손가락은 인장 모달리티가 잘 드러나는 장착(굽힘면)이다. 논문 기반 60A 벨로즈(공압)는 센서가 중립축이라 인장 변화가 작고 "
        "힘 용량이 약 6 N(200 kPa) — 하드웨어 선택 논의에 활용 (A–Z 보고서 부록 C).",
    ])
    D.h("6. 실물 실험 데이터 체크리스트")
    D.bl([
        "파지 시퀀스 로그 (1 kHz): t, ΔL, ΔR, V_TENG, gate_proba, ε̂, d̂/F̂, 상태, u, 재질 — 펌웨어 UART 출력에 상태·u·재질 열 추가.",
        "정답 계측: 물체 쪽 로드셀(힘), 스테이지 위치(근접), 손가락 각도(카메라 → 인장 환산).",
        "물체 세트: 금속 원통, 실리콘 손가락 모형(유전체), 계란(또는 파손 임계 알려진 시편), 스펀지.",
        "지연 실험: 같은 제어를 PC 경유(USB 스트리밍 → PC 추론 → 명령)로도 구현해 비교, 또는 펌웨어에 인위적 지연 버퍼.",
        "ablation: 펌웨어 플래그로 근접 스케줄·재질 판별·TENG 융합·안전 계층을 끄고 반복 (각 조건 ≥10회).",
        "도메인 적응: 그리퍼 장착 상태 보정 데이터(굽힘+누름 동시) N = 500…20 000, 0909 데이터와 혼합 fine-tune 후 재배포.",
    ])
    path = Path(out_dir) / GUIDE
    try:
        D.d.save(str(path))
    except PermissionError:
        path = Path(out_dir) / GUIDE.replace(".docx", "_갱신.docx")
        D.d.save(str(path))
    return path


def _f(x, nd=2):
    return "—" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x:.{nd}f}"


def _s1(n):
    return dict(
        msg="단일 전극 센서에서 액추에이터까지 전 과정이 MCU 하나 위 1 kHz 폐루프로 닫힌다.",
        panels="왼→오: 소프트 그리퍼(센서) → TDM 측정(L·V_TENG·R_DC 슬롯) → Gate-MoE 디커플링(게이트 ∨ TENG, 전문가 A/B) "
               "→ 다섯 모달리티 → 제어기(상태 선택·pre-shape·admittance·안전) → 액추에이터 → 물리 상호작용(피드백).",
        data="구조도. 수치(파라미터·용량·261 µs)는 X-CUBE-AI 리포트와 보드 로그 실측.",
        caption="Fig. 1. Embedded multimodal sensing-to-control pipeline. A single two-wire EGaIn electrode is "
                "time-division multiplexed into inductance, DC resistance and TENG voltage; an on-board gated "
                "mixture-of-experts decouples strain, proximity and force, fused with TENG contact and an "
                "inductance-residual material cue, and drives a state-machine controller at 1 kHz without a host PC.",
        body="그리퍼의 각 손가락 안쪽에 부착된 단일 전극 센서는 1 ms 주기로 L, R_DC, V_TENG를 시분할 측정한다. 게이트 MLP는 "
             "(ΔL, ΔR)로 접촉 여부를 판단하고, 접촉 신호는 TENG 펄스 래치와 OR로 융합된다. 비접촉에서는 전문가 A가 변형률과 근접 거리를, "
             "접촉 시에는 전문가 B가 변형률과 힘을 추정한다. 재질은 접근 중 ΔL에서 변형률 성분을 뺀 잔차의 부호로 판별한다. "
             "이 다섯 정보는 같은 MCU 위의 상태 머신 제어기로 바로 전달되어 액추에이터 명령을 만든다.",
        real="그대로 사용. 액추에이터·손가락 사양만 실제 하드웨어로 갱신.")


def _s2(n):
    return dict(
        msg=f"측정(≈200 µs)과 추론({INFER_US:.0f} µs)이 1 ms 안에 끝나고 절반 이상이 남는다 — 3지·4채널 확장 여유.",
        panels="왼쪽: 한 TDM 주기 안의 타이밍(LDC I²C-DMA, TENG/R ADC 창, MoE 추론, 제어, 유휴). 오른쪽: 네트워크별 파라미터·MACC·Flash·RAM.",
        data="측정 슬롯: 펌웨어 설정값(TENG/R 100 µs, LDC ~200 µs). 추론: 보드 로그 latency_us 중앙값/최대. 제어기 시간은 아직 "
             "보드 미구현이라 추정(<20 µs) — 실측으로 교체 필요.",
        caption=f"Fig. 2. Timing within one 1-ms TDM cycle and on-board model budget. Measurement slots run in parallel "
                f"via DMA; the full Gate-MoE inference takes {INFER_US:.0f} µs (median) on an STM32G473 at 170 MHz "
                f"(X-CUBE-AI, float32), leaving ≈50 % of the period idle.",
        body="게이트(68 파라미터)와 전문가 하나(674 파라미터)만 실행되므로 추론당 연산은 1,356 MAC이고, 가중치 5.5 KB·활성 메모리 160 B로 "
             "MCU 내부 메모리에 충분히 들어간다. 측정과 추론, 제어가 모두 TDM 주기 안에서 끝나므로 센서 데이터를 PC로 보내지 않고 "
             "같은 주기 안에서 제어 명령을 낼 수 있다.",
        real="제어기 펌웨어 구현 후 DWT 사이클 카운터로 '제어+명령' 구간 실측, 3지(3채널) 동시 동작 시 타이밍 재측정.")


def _s3(n):
    return dict(
        msg="상태 머신의 각 상태·전이가 어떤 모달리티로 움직이는지 한눈에.",
        panels="OPEN → APPROACH(근접으로 닫힘 속도, 잔차로 재질) → GRASP(재질별 목표 힘 admittance) → HOLD/LIFT → RELEASE. "
               "전이: 접근 명령, 세 손가락 접촉(게이트 ∨ TENG), 리프트·해제 명령, 3 s 타임아웃. 하단: 모든 상태에 걸친 안전 계층.",
        data="설계도 (sim_gripper/control/controller.py와 1:1).",
        caption="Fig. 3. Grasp state machine and the modality that drives each state. Proximity schedules the closing "
                "speed, the inductance residual latches the material class that sets the force target, gate/TENG "
                "contact triggers the transition to admittance force control, and a low-pass force/strain safety "
                "layer overrides every state.",
        body="Compliant 손가락에서는 힘을 직접 명령할 수 없으므로 접촉 전에는 위치(닫힘 속도)를, 접촉 후에는 힘 오차를 명령 속도로 "
             "바꾸는 admittance를 쓴다. 근접 거리 d̂에 따라 멀면 빠르게, 가까우면 느리게 닫아 접촉 충격을 줄이고, 접근 중 결정된 재질에 "
             "따라 목표 힘(도체 4 N, 그 외 1.5 N)을 고른다. 안전 계층은 저역통과한 힘·힘 변화율·변형률 범위를 감시해 어떤 상태에서도 "
             "명령을 개방 방향으로 되돌린다.",
        real="그대로 사용 (파라미터 값만 실측 보정 후 갱신).")


def _s4(n):
    return dict(
        msg="한 번의 파지(접근→접촉→파지→리프트→해제) 동안 다섯 모달리티가 모두 나타나고 제어 명령과 연결된다.",
        panels="a 근접(참/추정), b TENG 전압과 접촉 판정, c 힘(참/추정, F_t·F_max), d 인장(참/추정), e 원신호 ΔL·ΔR, "
               "f 명령 u와 재질 판정 시점. 배경 음영 = 상태.",
        data="시뮬레이션 E1 금속 원통, 판형 텐던 손가락, 그리퍼 보정 MoE, 손가락 0 (out/paper_runs/E1_metal_flat_*.pkl).",
        caption="Fig. 4. Simulated grasp of a metal cylinder showing all five modalities from one electrode: "
                "(a) proximity, (b) TENG contact, (c) force, (d) strain, (e) raw TDM signals and (f) actuator command "
                "with the latched material class. Shading marks controller states.",
        body="접근 단계에서 d̂가 줄어들며 닫힘 속도가 느려지고, 도체 근접으로 ΔL이 크게 감소하면서 재질이 '도체'로 확정된다. 접촉 순간 "
             "TENG 양의 펄스가 접촉 판정을 일으켜 admittance로 전환되고, 파지 중 ΔR 증가가 힘으로, 손가락 굽힘이 인장 감소로 분리된다. "
             "해제 시 TENG 음의 펄스가 분리를 확인한다.",
        real="같은 시퀀스를 실제 그리퍼로 촬영·로깅 (정답: 로드셀·스테이지·각도). 그림 코드 재사용 가능.")


def _s5(n):
    return dict(
        msg="접촉 전에 재질을 알아 목표 힘을 바꾼다 — 계란·손은 약하게, 금속은 강하게.",
        panels="a 접촉 직전 ΔL 잔차(도체 음, 유전체 양, 미정 띠 ±θ). b 접촉 후 실제 파지 힘과 재질별 목표(점선), 계란 파손선.",
        data="시뮬레이션 E1 금속·E2 손·E2 계란, 판형 텐던, 보정 MoE.",
        caption="Fig. 5. Pre-contact material recognition and material-adaptive force targets (simulated). (a) The sign "
                "of the inductance residual separates conductive and dielectric objects before contact; (b) the "
                "controller applies 4 N to the metal cylinder and 1.5 N to the finger and the egg, which stays intact.",
        body="ΔR은 비접촉에서 변형률에만 의존하므로 ΔR에서 역산한 변형률로 ΔL의 변형률 성분을 제거하면 근접 물체의 전기적 성질만 남는다. "
             "도체는 와전류로 L을 낮추고 유전체는 약간 높여, 잔차 부호가 50 ms 이상 유지되면 재질을 확정하고 첫 접촉 순간 고정한다.",
        real="실제 재질 세트(금속, 실리콘 손 모형, 계란, 플라스틱)로 판별 정확도 혼동행렬을 추가 (각 ≥20회).")


def _s6(n):
    lb, bel = n.get("lat_b1", {}), n.get("lat_bel", {})
    brk = n.get("lat_broken_b1", [])
    return dict(
        msg="센서→제어 지연이 수십 ms만 되어도 빠른 힘 제어는 계란을 깬다; 온보드 1 ms 루프는 높은 이득을 안전하게 허용한다.",
        panels="a 이득(B_a) × 지연 지도: 계란 최대 힘, ✕ = 파손 (판형 텐던). b B_a = 1에서 지연 대비 최대 힘 — 텐던(빠름) vs 공압 "
               "60A(느림).",
        data="시뮬레이션 E2 계란 4.5 s, 지연 = 센서 처리 결과를 FIFO로 늦춰 제어기에 전달 (fig6_latency_data.csv).",
        caption="Fig. 6. Benefit of closing the loop on board (simulated). Delaying the sensing result by the latency of "
                "a host-PC loop raises the peak egg force and breaks the egg at high control gains; the on-board "
                "1-ms loop keeps every gain safe. Slow pneumatic actuation masks the effect.",
        body=f"B_a = 1에서 계란 최대 힘은 지연 0 ms {_f(lb.get(0))} N → 50 ms {_f(lb.get(50))} N → 100 ms {_f(lb.get(100))} N"
             f"{', 파손 지연 ' + ', '.join(str(x) for x in brk) + ' ms' if brk else ''}. 빠른 텐던 구동일수록 지연에 민감하고, 느린 공압 "
             f"벨로즈는 지연 0 → 100 ms에서 {_f(bel.get(0))} → {_f(bel.get(100))} N으로 둔감하다. 즉 Embedded AI의 이점은 '빠른 "
             "액추에이터의 성능을 지연 없이 끌어낼 수 있다'는 것이다.",
        real="PC 경유 구현(USB 스트리밍 + PC 추론)과 온보드 구현을 같은 계란 시편으로 비교, 실제 지연 분포(평균·지터) 측정.")


def _s7(n):
    return dict(
        msg="경험식 해석적 역산은 근거리에서 발산하지만 Gate-MoE는 같은 영역에서 안정적이다.",
        panels="a 해석적 역산 발산률(ε–d), b 해석적 거리 오차(빨강 = 발산), c MoE 거리 오차, d ε 평균 곡선.",
        data="0909 실측 곡면 기반 신호 + 실측 노이즈, v4 MoE(평면 시험 도메인), 경험식은 연구 노트의 피팅식(R² 0.997).",
        caption="Fig. 7. Analytical inversion versus the embedded Gate-MoE on the flat-test domain. Newton inversion of "
                "the empirical L/R model diverges near contact (d < 4 mm), where the network keeps millimetre-level "
                "accuracy with a bounded output.",
        body=f"근거리(d < 4 mm)에서 해석적 역산은 발산률 {_f(100 * n.get('an_fail_near', np.nan), 0)} %인 반면 MoE의 거리 오차는 "
             f"{_f(n.get('moe_err_near'))} mm다. 중간·원거리에서는 두 방법 오차가 비슷하다(해석적 {_f(n.get('an_err_mid'))} mm, "
             f"MoE {_f(n.get('moe_err_mid'))} mm). 피팅식의 거듭제곱 항이 근거리 급변을 표현하지 못해 야코비안이 특이해지는 것이 원인이며, "
             "반복 계산이 필요 없는 MoE는 고정 시간(261 µs)에 끝난다.",
        real="같은 비교를 실측 테스트 세트(0909 test 창)로 직접 수행 가능 — 바로 본문용으로 교체 권장.")


def _s8(n):
    abl = n.get("abl")
    txt = ""
    if abl is not None:
        e_mat = _abl_row(abl, "− material", "E2_egg")
        h_mat = _abl_row(abl, "− material", "E2_hand")
        full_e, prox_e = _abl_row(abl, "full", "E2_egg"), _abl_row(abl, "− proximity", "E2_egg")
        txt = (f"재질 판별을 빼면 고정 목표 4 N 때문에 계란 최대 힘 {e_mat.peak:.2f} N"
               f"{'(파손)' if e_mat.broken else ''}, 손 {h_mat.peak:.2f} N(> 3 N). TENG 융합을 빼면 저하중·유전체 접촉을 못 잡아 "
               f"파지 자체가 실패한다. 근접 스케줄을 빼면 접촉 순간 힘 상승률이 {full_e.impact:.0f} → {prox_e.impact:.0f} N/s로 커진다. "
               "안전 계층은 정상 조건에서는 거의 개입하지 않는다(고장·외란 시나리오에서 역할 — 별도 실험 권장).")
    return dict(
        msg="모달리티를 하나씩 빼면 각자 다른 실패가 생긴다 — 다섯 모달리티가 모두 필요하다.",
        panels="a 조건별 최대 파지 힘과 성공(✓), 물체별 한계선. b 접촉 순간 최대 dF/dt (로그 축).",
        data="시뮬레이션 E2 계란·E2 손·E1 금속 × {full, −근접, −재질, −TENG, −안전} (fig8_ablation_data.csv).",
        caption="Fig. 8. Modality ablation (simulated). Removing material recognition overloads fragile objects, "
                "removing TENG contact fusion prevents contact detection, and removing proximity scheduling increases "
                "the impact force rate.",
        body=txt,
        real="펌웨어 플래그로 동일 ablation, 조건당 ≥10회 반복해 성공률·최대 힘 분포(박스플롯)로.")


def _s9(n):
    ds = n.get("ds")
    txt = ""
    if ds:
        last = list(ds["methods"]["mixed"])[-1]
        r = ds["methods"]["mixed"][last]
        cl = ds.get("closed_loop_E1_shift", {})
        ft = [v for k, v in cl.items() if k.startswith("finetune")]
        txt = (f"그리퍼 장착 시 평면 시험으로 학습한 v4의 변형률 R²는 {ds['v4_shift']['strain_r2']:.2f}까지 떨어진다. 같은 구조를 "
               f"N = {last} 보정 데이터와 평면 데이터를 섞어 fine-tune하면 그리퍼 ε/F R² {r['sim']['strain_r2']:.2f}/"
               f"{r['sim']['force_r2']:.2f}로 회복하고 평면 성능(ε R² {r['flat']['strain_r2']:.2f})도 유지한다."
               + (f" 폐루프에서는 최대 힘 {cl['v4']['max_force']:.1f} → {ft[0]['max_force']:.1f} N." if ft else ""))
    return dict(
        msg="그리퍼 장착에 따른 분포 변화를 같은 구조(재배포 가능)의 fine-tune으로 회복하며 평면 성능을 잊지 않는다.",
        panels="a 그리퍼 변형률 R², b 그리퍼 힘 R², c 평면 시험 변형률 R²(망각) — 보정 샘플 수 N 대비, 방법 4가지.",
        data="시뮬레이션 장착 편향(δε 3 %, γ 0.2, s 1.1) 주입 실험 (out/domain_shift.json). 게이트는 적응하지 않음(접촉은 TENG).",
        caption="Fig. 9. Recovering from the gripper-mounting domain shift (simulated). Mixed fine-tuning of the same "
                "1.4 k-parameter network restores gripper accuracy while preserving flat-test performance, so the "
                "adapted model can be redeployed on the MCU unchanged.",
        body=txt,
        real="실측 보정 데이터로 같은 곡선을 그리면 '필요한 보정 샘플 수'의 실험적 근거가 된다.")


_SPEC = {"fig1_architecture": _s1, "fig2_timing_budget": _s2, "fig3_state_machine": _s3,
         "fig4_grasp_timeline": _s4, "fig5_material_adaptive": _s5, "fig6_latency": _s6,
         "fig7_analytic_vs_moe": _s7, "fig8_ablation": _s8, "fig9_adaptation": _s9}


if __name__ == "__main__":
    print(build_guide())
